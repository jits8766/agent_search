#!/usr/bin/env python3
"""Global eval suite — HYBRID · EXPLORE · GUIDANCE · ANALYTICS (part of
``semantic_search.offline_harness``).

Runs the fixed ``QUERY_BANK`` against a live ``/search`` endpoint, scores each suite
(routing correctness, NDCG/P/R/Coh@5, SLA/timeout/error rates, L0/L1/L2 tier mix), and
writes to ``--out-dir`` (default ``auc-semantic-search/output/eval_search_queries/``, gitignored):

* ``eval_report_{stamp}.md`` — suite/section rollups, latency bands, misroutes
* ``eval_report_{stamp}.xlsx`` — ``queries`` + ``suite_summary`` + ``section_summary``
  + ``fail_only`` + ``misroutes`` + ``description``
* ``results_{stamp}.json`` — queries sheet (per-query; analysis-only source)
* sheet-wise JSON: ``suite_summary_{stamp}.json`` (includes holdout),
  ``section_summary_{stamp}.json``, ``fail_only_{stamp}.json``,
  ``misroutes_{stamp}.json`` — mirror xlsx sheets. No separate holdout_*.json.

Query concurrency (`_CONCURRENCY`) defaults from config/base.yaml's
``offline_eval.retrieval_eval.max_concurrent_queries`` (``EVAL_CONCURRENCY`` env
overrides); falls back to the empirically-tuned constant if config load fails.

Run directly (python one-liner, from auc-semantic-search/):
    python -m semantic_search.offline_harness.eval_test_search_queries http://localhost:8085
    python -m semantic_search.offline_harness.eval_test_search_queries --api http://localhost:8085 --suite HYBRID
    python -m semantic_search.offline_harness.eval_test_search_queries --api http://localhost:8085 --suite ANALYTICS
    python -m semantic_search.offline_harness.eval_test_search_queries --help

    # Rebuild sheet JSON + HOLDOUT on suite_summary xlsx from results_*.json — no HTTP:
    python -m semantic_search.offline_harness.eval_test_search_queries --analysis-only \\
        --results output/eval_search_queries/results_XXXX.json

Run over HTTP (see ``POST /internal/harness/eval-search`` in app.py):
    curl -sS -X POST http://localhost:8085/internal/harness/eval-search \\
        -F api=http://127.0.0.1:8085 -F suite=HYBRID
"""
import asyncio, hashlib, json, math, os, statistics, sys, time, tracemalloc, uuid
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

_SCRIPT_DIR = Path(__file__).resolve().parent
# offline_harness/ -> semantic_search/ -> semantic-search/ -> packages/ -> auc-semantic-search/
_REPO_ROOT = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(_REPO_ROOT / "packages" / "semantic-search"))
from semantic_search.config.loader import load_config  # noqa: E402
from semantic_search.config.models import AgentSearchConfig  # noqa: E402

_MAX_QUERY_CONCURRENCY = 4  # hard cap regardless of config/env override
# Defaults match base.yaml when config load fails.
_CONFIG_MAX_CONCURRENCY = _MAX_QUERY_CONCURRENCY
_REL_THRESHOLD: float = 0.7
try:
    # Prefer raw YAML keys so a partial/invalid typed config (e.g. missing seed DB
    # env) cannot block reading retrieval.metrics.relevance_threshold.
    _raw_cfg = load_config()
    if isinstance(_raw_cfg, dict):
        _re = (_raw_cfg.get("offline_eval") or {}).get("retrieval_eval") or {}
        if _re.get("max_concurrent_queries") is not None:
            _CONFIG_MAX_CONCURRENCY = int(_re["max_concurrent_queries"])
        _rm = (_raw_cfg.get("retrieval") or {}).get("metrics") or {}
        if _rm.get("relevance_threshold") is not None:
            _REL_THRESHOLD = float(_rm["relevance_threshold"])
    else:
        _cfg = AgentSearchConfig.from_dict(_raw_cfg)
        _CONFIG_MAX_CONCURRENCY = _cfg.offline_eval.retrieval_eval.max_concurrent_queries
        _REL_THRESHOLD = float(_cfg.retrieval.metrics.relevance_threshold)
except Exception:  # noqa: BLE001 - config-sourced defaults only; hard caps below still apply
    pass
if not (0.0 <= float(_REL_THRESHOLD) <= 1.0):
    _REL_THRESHOLD = 0.7

DEFAULT_OUTPUT_DIR = _REPO_ROOT / "output" / "eval_search_queries"

# ── Run configuration ─────────────────────────────────────────────────────────
_CONCURRENCY: int = max(   # concurrent requests per API. 4 keeps the local single-node
    1,                      # ClickHouse below saturation on the SUSTAINED full run (a 92-query
    min(                    # burst tolerates 6, but 696 queries sustained does not). The
        _MAX_QUERY_CONCURRENCY,  # middleware fix means overload degrades to graceful 503/fallback,
        int(os.environ.get("EVAL_CONCURRENCY", str(_CONFIG_MAX_CONCURRENCY))),
    ),                       # never a 500. Higher throughput WITH good results needs CH
)                            # capacity (container CPU/mem, cached price-bands + explore rails,
                             # or a CH cluster), not eval changes.
_TOP_K: int           = 50       # top_k sent to /search (API default); eval computed at @5 locally
_TIMEOUT_S: float     = 25.0     # max seconds per query (hard cap)
# Relevance proxy threshold: same key as /search retrieval_metrics
# (retrieval.metrics.relevance_threshold in base.yaml). Scores >= threshold count
# as relevant for local P/R/F1 (no gold labels).
_DEFAULT_SLA_MS: int  = 20000    # fallback SLA; per-suite SLA in SUITE_SLA_MS takes precedence

# ── Routing expectations per suite ───────────────────────────────────────────
# answer_mode values that count as "correctly routed" for each suite
EXPECTED_MODE: Dict[str, Set[str]] = {
    "HYBRID":    {"search"},
    "EXPLORE":   {"explore_fallback", "explore"},
    "GUIDANCE":  {"guidance"},
    "ANALYTICS": {"analytics"},
}

# Per-suite SLA thresholds — independent of --sla-ms CLI arg (used for fallback only)
SUITE_SLA_MS: Dict[str, int] = {
    "HYBRID":    10000,
    "EXPLORE":   10000,
    "GUIDANCE":  10000,
    "ANALYTICS": 20000,
}

# Latency product bands (reporting only — does not change suite SLA above).
LATENCY_LE_4S_MS: int = 4000
LATENCY_LE_10S_MS: int = 10000
# HYBRID p95 product target: OK if p95<=4s; WARN if p95<=suite SLA (10s); else BAD.
HYBRID_P95_TARGET_MS: int = LATENCY_LE_4S_MS
HYBRID_P95_WARN_MS: int = SUITE_SLA_MS["HYBRID"]

# failure_mode values that are expected/graceful — system working as designed.
# inventory_empty* = no matching domains in marketplace (not a bug).
# *_fallback / filter_conflict = guard/fallback rails fired intentionally.
# These are NOT counted as "true failures" in the report.
_EXPECTED_FALLBACKS: frozenset = frozenset({
    "inventory_empty",
    "inventory_empty_under_filters",
    "find_fallback",
    "explore_fallback_rail",
    "analytics_failure_fallback",
    "filter_conflict",
    "hybrid_explore_complement",
    "force_semantic_nonempty",
})

# Benchmark targets per suite per metric.
# Tuple: (warn_threshold, good_threshold, direction, note)
# direction='hi' -> higher is better (OK if >=good, WARN if >=warn, BAD otherwise)
# direction='lo' -> lower is better (OK if <=good, WARN if <=warn, BAD otherwise)
SUITE_BENCHMARKS: Dict[str, Dict[str, tuple]] = {
    "HYBRID": {
        "routing":    (70, 90, "hi", "lower=misroutes to wrong mode"),
        "entity_rate":(40, 55, "hi", "lower=filter coverage gap; higher=better extraction"),
        # lat_p95 uses _lat_p95_grade (ms thresholds), not _grade (%).
        "lat_p95_ms": (HYBRID_P95_WARN_MS, HYBRID_P95_TARGET_MS, "lo", "product target 4s"),
    },
    "EXPLORE": {
        "routing":    (70, 90, "hi", "lower=misroutes to wrong mode"),
        "entity_rate":(30, 15, "lo", "EXPLORE has no filters; high rate=false-positive entities"),
    },
    "GUIDANCE": {
        "routing":    (70, 90, "hi", "lower=misroutes to wrong mode"),
        "entity_rate":(25, 15, "lo", "GUIDANCE has no filters; high rate=false-positive entities"),
    },
    "ANALYTICS": {
        "routing":    (70, 90, "hi", "lower=misroutes to wrong mode"),
        "entity_rate":(35, 20, "lo", "ANALYTICS rarely has filters; high rate=false-positive entities"),
    },
}

# Decision-tier labels the QI ensemble actually emits today. The resolver
# stamps a tier by the winning voter (entity / semantic / llm); the engine adds
# multi_intent + the fallback labels. The old graduated cascade (L3 decompose /
# L4 assist / *_llm / *_entity compounds) no longer exists — those labels were
# removed here so the table reflects produced tiers only.
_L0_TIERS: frozenset = frozenset({
    "L0_entity", "L0_multi_intent",
})
_L1_TIERS: frozenset = frozenset({
    "L1_semantic",
})
_L2_TIERS: frozenset = frozenset({
    "L2_llm",
})
# Non-routing outcomes — abstain / timeout / no-slice / resolver error. Bucketed
# explicitly so they stop hiding in the gap between L0+L1+L2 and nv.
_FALLBACK_TIERS: frozenset = frozenset({
    "fallback", "L0_fallback", "ensemble_all_abstain",
})
_ALL_KNOWN_TIERS: frozenset = _L0_TIERS | _L1_TIERS | _L2_TIERS | _FALLBACK_TIERS

# ── Query bank ────────────────────────────────────────────────────────────────
QUERY_BANK: List[Tuple[str, str, str]] = [
    # ── HYBRID ───────────────────────────────────────────────────────────────
    # Semantic + Filters
    ("HYBRID", "Semantic + Filters", "ai startup domain under 1500 with some traffic"),
    ("HYBRID", "Semantic + Filters", "premium ai domain with real visitors"),
    ("HYBRID", "Semantic + Filters", "b2b saas startup name good branding"),
    ("HYBRID", "Semantic + Filters", "fintech startup domain under 2k no crypto vibe"),
    ("HYBRID", "Semantic + Filters", "healthcare startup name trustworthy"),
    ("HYBRID", "Semantic + Filters", "cybersecurity startup domain under 1k"),
    ("HYBRID", "Semantic + Filters", "climate tech domains added recently"),
    ("HYBRID", "Semantic + Filters", "ecommerce domains with existing traffic"),
    ("HYBRID", "Semantic + Filters", "devtools startup domain ending soon"),
    ("HYBRID", "Semantic + Filters", "cloud platform domain traffic 1000+"),
    ("HYBRID", "Semantic + Filters", "edtech domain under 1k easy to spell"),
    ("HYBRID", "Semantic + Filters", "travel startup name with real visitors"),
    ("HYBRID", "Semantic + Filters", "logistics saas domain good authority cheap"),
    ("HYBRID", "Semantic + Filters", "food delivery brand domain under 800"),
    ("HYBRID", "Semantic + Filters", "real estate domain decent traffic mid budget"),
    ("HYBRID", "Semantic + Filters", "hr software domain short and clean"),
    ("HYBRID", "Semantic + Filters", "gaming startup domain under 2000 catchy"),
    ("HYBRID", "Semantic + Filters", "legal tech domain trustworthy under 1500"),
    ("HYBRID", "Semantic + Filters", "looking for b2b saas or maybe fintech domain under 2k but not sure which direction my startup is going yet"),
    ("HYBRID", "Semantic + Filters", "something in tech space that has real traffic below 1500 or 2000 if backlinks are good"),
    # Similarity Search
    ("HYBRID", "Similarity Search", "domains like stripe or plaid"),
    ("HYBRID", "Similarity Search", "names with notion kinda vibe"),
    ("HYBRID", "Similarity Search", "similar to linear app style naming"),
    ("HYBRID", "Similarity Search", "domain like figma but not design related"),
    ("HYBRID", "Similarity Search", "cheaper alternative to databricks domain"),
    ("HYBRID", "Similarity Search", "openai style name no ai or gpt words"),
    ("HYBRID", "Similarity Search", "startup names like slack stripe zoom"),
    ("HYBRID", "Similarity Search", "yc startup type domain"),
    ("HYBRID", "Similarity Search", "something that sounds like vercel"),
    ("HYBRID", "Similarity Search", "names similar to airbnb feel"),
    ("HYBRID", "Similarity Search", "domain vibe like canva but cheaper"),
    ("HYBRID", "Similarity Search", "short names like uber or lyft"),
    ("HYBRID", "Similarity Search", "sounds like shopify but for services"),
    ("HYBRID", "Similarity Search", "names in the spirit of dropbox"),
    ("HYBRID", "Similarity Search", "i want something like stripe or notion or maybe figma not sure what vibe but clean and short and cheap"),
    ("HYBRID", "Similarity Search", "want a name that sounds like a startup you would see on product hunt or yc demo day kinda vibe"),
    # Brandability
    ("HYBRID", "Brandability", "short catchy .io names below 500"),
    ("HYBRID", "Brandability", "premium sounding one word domains"),
    ("HYBRID", "Brandability", "memorable enterprise saas domain"),
    ("HYBRID", "Brandability", "startup sounding domain under 500"),
    ("HYBRID", "Brandability", "trust evoking startup names"),
    ("HYBRID", "Brandability", "one word brandable domain under 3k"),
    ("HYBRID", "Brandability", "luxury sounding brand under budget"),
    ("HYBRID", "Brandability", "strong brand no traffic data needed"),
    ("HYBRID", "Brandability", "easy to say domain under 1k"),
    ("HYBRID", "Brandability", "punchy two syllable brand name"),
    ("HYBRID", "Brandability", "modern sounding tech brand cheap"),
    ("HYBRID", "Brandability", "clean minimal brand domain under 700"),
    ("HYBRID", "Brandability", "need short catchy domain sounds premium but also dont want to spend too much maybe 500 or under"),
    ("HYBRID", "Brandability", "one word or two word domain sounds clean and modern and maybe has some backlinks or authority too"),
    # Brand / Style Discovery
    ("HYBRID", "Brand / Style Discovery", "cool startup names available now"),
    ("HYBRID", "Brand / Style Discovery", "brandable domains worth browsing"),
    ("HYBRID", "Brand / Style Discovery", "clean b2b saas names"),
    ("HYBRID", "Brand / Style Discovery", "premium startup brands available"),
    ("HYBRID", "Brand / Style Discovery", "fresh modern names worth a look"),
    ("HYBRID", "Brand / Style Discovery", "understated luxury brand domains"),
    ("HYBRID", "Brand / Style Discovery", "show me cool startup names available and also look like they could be real company someday"),
    ("HYBRID", "Brand / Style Discovery", "browse names that could work for any b2b startup or tech product and also look good on business card"),
    ("HYBRID", "Brand / Style Discovery", "show me modern startup names that are likely domain-available, memorable after hearing once, suitable for venture-backed B2B software, and distinct enough to avoid confusion with existing tech companies"),
    # Category Filter Search
    ("HYBRID", "Category Filter Search", "show ai domains available now"),
    ("HYBRID", "Category Filter Search", "fintech domains worth browsing"),
    ("HYBRID", "Category Filter Search", "healthcare domains available"),
    ("HYBRID", "Category Filter Search", "cyber security names"),
    ("HYBRID", "Category Filter Search", "devtools domains right now"),
    ("HYBRID", "Category Filter Search", "crypto domains available today"),
    ("HYBRID", "Category Filter Search", "saas domains under 1k"),
    ("HYBRID", "Category Filter Search", "green energy domains available"),
    ("HYBRID", "Category Filter Search", "ai agent domains right now"),
    ("HYBRID", "Category Filter Search", "biotech names available"),
    ("HYBRID", "Category Filter Search", "show me ai domains but also maybe healthtech or fintech ones too if they look good"),
    ("HYBRID", "Category Filter Search", "browse ai and fintech domains together i cant decide which category my startup fits into"),
    # SEO + Authority
    ("HYBRID", "SEO + Authority", "domains with strong backlinks under 2k"),
    ("HYBRID", "SEO + Authority", "good seo domains cheap"),
    ("HYBRID", "SEO + Authority", "traffic rich domains under 3k"),
    ("HYBRID", "SEO + Authority", "strong authority domains for investing"),
    ("HYBRID", "SEO + Authority", "old domains with backlink juice"),
    ("HYBRID", "SEO + Authority", "active websites traffic 5k+ under 3k"),
    ("HYBRID", "SEO + Authority", "undervalued domains with authority"),
    ("HYBRID", "SEO + Authority", "seo value domains expiring now"),
    ("HYBRID", "SEO + Authority", "high da domains under 1500"),
    ("HYBRID", "SEO + Authority", "aged domains with clean backlinks"),
    ("HYBRID", "SEO + Authority", "domains with plenty referring domains cheap"),
    ("HYBRID", "SEO + Authority", "blog domains with steady traffic under 2k"),
    ("HYBRID", "SEO + Authority", "need domain with good backlinks or high da under 2k but also should sound like real brand"),
    ("HYBRID", "SEO + Authority", "want strong seo domain under 1500 with decent da and also should be short and brandable ideally"),
    # Expired + Auctions
    ("HYBRID", "Expired + Auctions", "expired one word .com under 1k"),
    ("HYBRID", "Expired + Auctions", "expiring domains decent traffic cheap"),
    ("HYBRID", "Expired + Auctions", "backorder worthy domains"),
    ("HYBRID", "Expired + Auctions", "old expired .com with traffic"),
    ("HYBRID", "Expired + Auctions", "startup sounding domains ending soon"),
    ("HYBRID", "Expired + Auctions", "buy now expired domains only"),
    ("HYBRID", "Expired + Auctions", "expired ai domains under 500"),
    ("HYBRID", "Expired + Auctions", "dropping .com with backlinks cheap"),
    ("HYBRID", "Expired + Auctions", "auction domains ending tonight under 1k"),
    ("HYBRID", "Expired + Auctions", "recently expired brandable names"),
    ("HYBRID", "Expired + Auctions", "expired domain or auction ending soon doesnt matter which just needs traffic or backlinks and cheap"),
    ("HYBRID", "Expired + Auctions", "find me expired or dropping domain that has traffic and also sounds like real brand not keyword spam"),
    # Structured Filters
    ("HYBRID", "Structured Filters", ".net domains under 100"),
    ("HYBRID", "Structured Filters", ".org domains under 200"),
    ("HYBRID", "Structured Filters", "short .com less than 8 chars"),
    ("HYBRID", "Structured Filters", "four letter .com no numbers"),
    ("HYBRID", "Structured Filters", "one word domain any extension"),
    ("HYBRID", "Structured Filters", "cloud keyword .io domains"),
    ("HYBRID", "Structured Filters", "no hyphen .com domains under 300"),
    ("HYBRID", "Structured Filters", "domains underpriced compared to similar sales"),
    ("HYBRID", "Structured Filters", "five letter .com no hyphens"),
    ("HYBRID", "Structured Filters", ".ai domains under 2000"),
    ("HYBRID", "Structured Filters", "three letter domains any tld"),
    ("HYBRID", "Structured Filters", "domains with no numbers under 500"),
    ("HYBRID", "Structured Filters", ".co domains short and cheap"),
    ("HYBRID", "Structured Filters", "dictionary word .com under 5k"),
    ("HYBRID", "Structured Filters", "short .com no numbers and no hyphens under 500 but also should sound like it could be real startup"),
    ("HYBRID", "Structured Filters", "need domain thats either .com or .io short no numbers no hyphens and under 1k preferably"),
    # Investor Semantic
    ("HYBRID", "Investor Semantic", "domains with good flip potential"),
    ("HYBRID", "Investor Semantic", "what domains investors buying lately"),
    ("HYBRID", "Investor Semantic", "domains likely worth more in future"),
    ("HYBRID", "Investor Semantic", "domains with strong resale value"),
    ("HYBRID", "Investor Semantic", "hidden gem domains investors might buy"),
    ("HYBRID", "Investor Semantic", "undervalued names with upside"),
    ("HYBRID", "Investor Semantic", "good long term hold domains"),
    ("HYBRID", "Investor Semantic", "investor grade domains under 5k"),
    ("HYBRID", "Investor Semantic", "domains with high liquidity potential"),
    ("HYBRID", "Investor Semantic", "strong portfolio addition names"),
    ("HYBRID", "Investor Semantic", "quick flip domains under 1k"),
    ("HYBRID", "Investor Semantic", "find me domains with flip potential but also should have some traffic or backlinks or authority"),
    ("HYBRID", "Investor Semantic", "want domains that appreciate in value over time but also could work as actual business name right now"),
    # Value / Hidden Gems
    ("HYBRID", "Value / Hidden Gems", "hidden gem domains"),
    ("HYBRID", "Value / Hidden Gems", "underpriced domains worth checking"),
    ("HYBRID", "Value / Hidden Gems", "good value names right now"),
    ("HYBRID", "Value / Hidden Gems", "overlooked domains with potential"),
    ("HYBRID", "Value / Hidden Gems", "sleeper picks available now"),
    ("HYBRID", "Value / Hidden Gems", "domains nobody noticed yet"),
    ("HYBRID", "Value / Hidden Gems", "cheap names that look expensive"),
    ("HYBRID", "Value / Hidden Gems", "bargain brandables under 300"),
    ("HYBRID", "Value / Hidden Gems", "hidden gems that are cheap but also look expensive and have some traffic or backlinks too"),
    ("HYBRID", "Value / Hidden Gems", "sleeper picks nobody knows about but should also be short and pronounceable and not too weird"),
    # Ambiguous / Open-Ended Hybrid
    ("HYBRID", "Ambiguous / Open-Ended Hybrid", "what would you buy if investing today"),
    ("HYBRID", "Ambiguous / Open-Ended Hybrid", "show me sleeper domains"),
    ("HYBRID", "Ambiguous / Open-Ended Hybrid", "domains flying under radar"),
    ("HYBRID", "Ambiguous / Open-Ended Hybrid", "any overlooked premium names"),
    ("HYBRID", "Ambiguous / Open-Ended Hybrid", "good value compared to market"),
    ("HYBRID", "Ambiguous / Open-Ended Hybrid", "bargain domains with upside"),
    ("HYBRID", "Ambiguous / Open-Ended Hybrid", "future category winners"),
    ("HYBRID", "Ambiguous / Open-Ended Hybrid", "domains in growing sectors"),
    ("HYBRID", "Ambiguous / Open-Ended Hybrid", "opportunities nobody noticed yet"),
    ("HYBRID", "Ambiguous / Open-Ended Hybrid", "best value premium domains right now"),
    ("HYBRID", "Ambiguous / Open-Ended Hybrid", "surprise me with something undervalued"),
    ("HYBRID", "Ambiguous / Open-Ended Hybrid", "what looks cheap but isnt"),
    ("HYBRID", "Ambiguous / Open-Ended Hybrid", "just show me anything good under 2k that has traffic or backlinks and sounds like it could be startup"),
    ("HYBRID", "Ambiguous / Open-Ended Hybrid", "not sure if i want expired or brandable or what but show me good options under 1k with any kind of value"),
    # Messy Real User Queries
    ("HYBRID", "Messy Real User Queries", "good domain for flipping"),
    ("HYBRID", "Messy Real User Queries", "what investor buying now"),
    ("HYBRID", "Messy Real User Queries", "any undervalued domains today"),
    ("HYBRID", "Messy Real User Queries", "show me hidden gems for resale"),
    ("HYBRID", "Messy Real User Queries", "domains worth investing in rn"),
    ("HYBRID", "Messy Real User Queries", "what names can sell later higher"),
    ("HYBRID", "Messy Real User Queries", "cheap premium domains maybe"),
    ("HYBRID", "Messy Real User Queries", "any sleeper picks"),
    ("HYBRID", "Messy Real User Queries", "domains below market value"),
    ("HYBRID", "Messy Real User Queries", "what would domain pro buy"),
    ("HYBRID", "Messy Real User Queries", "lookin for cheap ai domain that sounds legit"),
    ("HYBRID", "Messy Real User Queries", "need a short brandable name not too pricey"),
    ("HYBRID", "Messy Real User Queries", "find me somethin like a fintech brand under 2k"),
    ("HYBRID", "Messy Real User Queries", "want a .com thats short no numbers cheap"),
    ("HYBRID", "Messy Real User Queries", "need domain for my startup its like ai meets fintech not sure which way to go budget around 2k maybe less"),
    ("HYBRID", "Messy Real User Queries", "want short name sounds legit probs .com or .io under 1k but if its expired with backlinks ill pay more"),
    # Misspelled / Typo Queries
    ("HYBRID", "Misspelled / Typo Queries", "ai startp domain under 1500"),
    ("HYBRID", "Misspelled / Typo Queries", "premiun doamin with trafic"),
    ("HYBRID", "Misspelled / Typo Queries", "fintec startup name cheep"),
    ("HYBRID", "Misspelled / Typo Queries", "healthcar domain trustworthy"),
    ("HYBRID", "Misspelled / Typo Queries", "expird .com with backlinks"),
    ("HYBRID", "Misspelled / Typo Queries", "brandible name under 500"),
    ("HYBRID", "Misspelled / Typo Queries", "short cathy .io domain"),
    ("HYBRID", "Misspelled / Typo Queries", "domian like stripe"),
    ("HYBRID", "Misspelled / Typo Queries", "availabe ai domains now"),
    ("HYBRID", "Misspelled / Typo Queries", "devtols domain ending soon"),
    ("HYBRID", "Misspelled / Typo Queries", "cybersecuirty domain under 1k"),
    ("HYBRID", "Misspelled / Typo Queries", "undervaluedd domains today"),
    ("HYBRID", "Misspelled / Typo Queries", "saas startp name good brandng"),
    ("HYBRID", "Misspelled / Typo Queries", "trafic rich domain cheap"),
    ("HYBRID", "Misspelled / Typo Queries", "ai or fintec domian under 2k wit trafic and goood brandng"),
    ("HYBRID", "Misspelled / Typo Queries", "short .com no numbrs no hyphn under 1k brandible for tech compny"),
    # Incomplete / Fragment Queries
    ("HYBRID", "Incomplete / Fragment Queries", "ai domain under"),
    ("HYBRID", "Incomplete / Fragment Queries", "short .com no"),
    ("HYBRID", "Incomplete / Fragment Queries", "something fintech cheap"),
    ("HYBRID", "Incomplete / Fragment Queries", "expired with traffic"),
    ("HYBRID", "Incomplete / Fragment Queries", "brandable like"),
    ("HYBRID", "Incomplete / Fragment Queries", "healthcare name trust"),
    ("HYBRID", "Incomplete / Fragment Queries", "one word .io"),
    ("HYBRID", "Incomplete / Fragment Queries", "cheap good seo"),
    ("HYBRID", "Incomplete / Fragment Queries", "domain investor flip"),
    ("HYBRID", "Incomplete / Fragment Queries", "premium under 1k"),
    ("HYBRID", "Incomplete / Fragment Queries", "hidden gem"),
    ("HYBRID", "Incomplete / Fragment Queries", "ending soon traffic"),
    ("HYBRID", "Incomplete / Fragment Queries", "ai domain under 2k with"),
    ("HYBRID", "Incomplete / Fragment Queries", "something like stripe but"),
    # Long Natural Sentences
    ("HYBRID", "Long Natural Sentences", "i am starting an ai company and need a short domain under 1500 that has some traffic already"),
    ("HYBRID", "Long Natural Sentences", "looking for a fintech brand name that doesnt sound like crypto and stays under 2 thousand"),
    ("HYBRID", "Long Natural Sentences", "can you find me a domain similar to stripe but cheaper and with no numbers in it"),
    ("HYBRID", "Long Natural Sentences", "i want something brandable for my saas under 500 dollars that is easy to remember"),
    ("HYBRID", "Long Natural Sentences", "need an expired .com with decent backlinks for seo but i dont want to spend over 2k"),
    ("HYBRID", "Long Natural Sentences", "find a healthcare domain that sounds trustworthy and is available to buy right now"),
    ("HYBRID", "Long Natural Sentences", "show me one word .io domains that sound premium but are still under a thousand"),
    ("HYBRID", "Long Natural Sentences", "i need a domain for my devtools startup that ends in auction soon and is cheap"),
    ("HYBRID", "Long Natural Sentences", "looking for an undervalued domain investors would grab with good resale potential"),
    ("HYBRID", "Long Natural Sentences", "help me find a clean b2b saas name with good branding and some existing traffic"),
    ("HYBRID", "Long Natural Sentences", "want a domain like notion vibe but not too expensive and short enough to type"),
    ("HYBRID", "Long Natural Sentences", "find me a cybersecurity startup name under 1000 that looks professional"),
    ("HYBRID", "Long Natural Sentences", "i would like an aged domain with real backlinks under 1500 for ranking purposes"),
    ("HYBRID", "Long Natural Sentences", "show domains in growing sectors that look cheap now but could be worth more later"),
    ("HYBRID", "Long Natural Sentences", "i am looking for domain that works for either ai startup or fintech product and has some traffic already and stays under 2000 dollars"),
    ("HYBRID", "Long Natural Sentences", "looking for domain that could work as real company name but also has good resale value in case startup doesnt work out under 3k"),
    # Price-Range Specific
    ("HYBRID", "Price-Range Specific", "domains between 500 and 1000"),
    ("HYBRID", "Price-Range Specific", "ai domains 1k to 2k with traffic"),
    ("HYBRID", "Price-Range Specific", "anything under 250 brandable"),
    ("HYBRID", "Price-Range Specific", "premium names 3k to 5k worth it"),
    ("HYBRID", "Price-Range Specific", "cheap domains below 100"),
    ("HYBRID", "Price-Range Specific", "mid budget saas names around 1500"),
    ("HYBRID", "Price-Range Specific", "domains exactly under 750"),
    ("HYBRID", "Price-Range Specific", "one word .com under 10k"),
    ("HYBRID", "Price-Range Specific", "fintech names 800 to 1200"),
    ("HYBRID", "Price-Range Specific", "expired domains under 300 with traffic"),
    ("HYBRID", "Price-Range Specific", ".io domains capped at 500"),
    ("HYBRID", "Price-Range Specific", "brandables no more than 2k"),
    ("HYBRID", "Price-Range Specific", "domains around 1k with backlinks"),
    ("HYBRID", "Price-Range Specific", "budget friendly .com under 400"),
    ("HYBRID", "Price-Range Specific", "anything between 500 and 1500 that has traffic and also sounds like real brand not just keyword"),
    ("HYBRID", "Price-Range Specific", "cheap domains under 300 but also should have some authority or at least be really brandable and short"),
    # TLD + Niche Combos
    ("HYBRID", "TLD + Niche Combos", "ai .com domains with traffic"),
    ("HYBRID", "TLD + Niche Combos", "fintech .io names under 1k"),
    ("HYBRID", "TLD + Niche Combos", "health .co domains brandable"),
    ("HYBRID", "TLD + Niche Combos", "crypto .xyz domains cheap"),
    ("HYBRID", "TLD + Niche Combos", "saas .app names available"),
    ("HYBRID", "TLD + Niche Combos", "dev .io domains short"),
    ("HYBRID", "TLD + Niche Combos", "cloud .com under 3k"),
    ("HYBRID", "TLD + Niche Combos", "shop .store domains available"),
    ("HYBRID", "TLD + Niche Combos", "legal .law domains trustworthy"),
    ("HYBRID", "TLD + Niche Combos", "travel .com with backlinks"),
    ("HYBRID", "TLD + Niche Combos", "ai .com or .io doesnt matter under 1k with some traffic and good branding for tech startup"),
    ("HYBRID", "TLD + Niche Combos", "saas .io or .app names short under 500 but also should sound modern and professional"),

    # ── EXPLORE ──────────────────────────────────────────────────────────────
    # Trending / What's Hot
    ("EXPLORE", "Trending / What's Hot", "whats trending in domains right now"),
    ("EXPLORE", "Trending / What's Hot", "show me hot domains this week"),
    ("EXPLORE", "Trending / What's Hot", "what people looking at most"),
    ("EXPLORE", "Trending / What's Hot", "domains getting lots of attention"),
    ("EXPLORE", "Trending / What's Hot", "hottest categories right now"),
    ("EXPLORE", "Trending / What's Hot", "anything interesting happening today"),
    ("EXPLORE", "Trending / What's Hot", "trending ai domains this week"),
    ("EXPLORE", "Trending / What's Hot", "whats buzzing in the marketplace"),
    ("EXPLORE", "Trending / What's Hot", "hot picks everyone talking about"),
    ("EXPLORE", "Trending / What's Hot", "domains heating up lately"),
    ("EXPLORE", "Trending / What's Hot", "trending tlds right now"),
    ("EXPLORE", "Trending / What's Hot", "whats popular in fintech domains"),
    ("EXPLORE", "Trending / What's Hot", "show me whats trending in ai and fintech domains together this week i want both"),
    ("EXPLORE", "Trending / What's Hot", "trending domains this week but also most viewed because i want to see both popularity signals"),
    # Fresh Listings
    ("EXPLORE", "Fresh Listings", "new domains added today"),
    ("EXPLORE", "Fresh Listings", "latest premium listings"),
    ("EXPLORE", "Fresh Listings", "fresh domains worth checking"),
    ("EXPLORE", "Fresh Listings", "what got listed this week"),
    ("EXPLORE", "Fresh Listings", "newest startup domains"),
    ("EXPLORE", "Fresh Listings", "just listed domains right now"),
    ("EXPLORE", "Fresh Listings", "recently added .com domains"),
    ("EXPLORE", "Fresh Listings", "fresh ai names this week"),
    ("EXPLORE", "Fresh Listings", "new arrivals worth a look"),
    ("EXPLORE", "Fresh Listings", "latest drops added today"),
    ("EXPLORE", "Fresh Listings", "brand new brandable names"),
    ("EXPLORE", "Fresh Listings", "whats new in the marketplace"),
    ("EXPLORE", "Fresh Listings", "new domains added today but also show me if any of them have traffic or backlinks already"),
    ("EXPLORE", "Fresh Listings", "fresh listings this week but also want to see if any are ending soon so i can act fast"),
    # Ending Soon
    ("EXPLORE", "Ending Soon", "domains ending today"),
    ("EXPLORE", "Ending Soon", "auctions ending next few hours"),
    ("EXPLORE", "Ending Soon", "ending this weekend worth watching"),
    ("EXPLORE", "Ending Soon", "last chance domains"),
    ("EXPLORE", "Ending Soon", "good auctions ending soon"),
    ("EXPLORE", "Ending Soon", "domains ending next 24 hrs"),
    ("EXPLORE", "Ending Soon", "pending delete domains worth grabbing"),
    ("EXPLORE", "Ending Soon", "closing tonight worth watching"),
    ("EXPLORE", "Ending Soon", "auctions wrapping up soon"),
    ("EXPLORE", "Ending Soon", "final hours domains"),
    ("EXPLORE", "Ending Soon", "ending tomorrow morning"),
    ("EXPLORE", "Ending Soon", "soon to close brandables"),
    ("EXPLORE", "Ending Soon", "domains ending soon but also show me only ones with traffic or backlinks or some real value"),
    ("EXPLORE", "Ending Soon", "last chance domains but also specifically in ai or fintech categories if possible"),
    # Popular / Social Proof
    ("EXPLORE", "Popular / Social Proof", "most viewed domains"),
    ("EXPLORE", "Popular / Social Proof", "domains getting most bids"),
    ("EXPLORE", "Popular / Social Proof", "most active auctions"),
    ("EXPLORE", "Popular / Social Proof", "what everyone bidding on"),
    ("EXPLORE", "Popular / Social Proof", "top watched domains"),
    ("EXPLORE", "Popular / Social Proof", "popular .ai domains lately"),
    ("EXPLORE", "Popular / Social Proof", "most favorited domains"),
    ("EXPLORE", "Popular / Social Proof", "domains with most watchers"),
    ("EXPLORE", "Popular / Social Proof", "crowd favorite names"),
    ("EXPLORE", "Popular / Social Proof", "what people adding to watchlist"),
    ("EXPLORE", "Popular / Social Proof", "busiest auctions today"),
    ("EXPLORE", "Popular / Social Proof", "most clicked domains this week"),
    ("EXPLORE", "Popular / Social Proof", "most viewed domains this week but also show me if any are ending soon so i can decide fast"),
    ("EXPLORE", "Popular / Social Proof", "crowd favorites but also want to see something new that hasnt gotten attention yet too"),
    # New Categories / Niches Browse
    ("EXPLORE", "New Categories / Niches Browse", "browse ai domains"),
    ("EXPLORE", "New Categories / Niches Browse", "show me fintech listings"),
    ("EXPLORE", "New Categories / Niches Browse", "explore healthcare names"),
    ("EXPLORE", "New Categories / Niches Browse", "look through crypto domains"),
    ("EXPLORE", "New Categories / Niches Browse", "browse short .com names"),
    ("EXPLORE", "New Categories / Niches Browse", "explore one word domains"),
    ("EXPLORE", "New Categories / Niches Browse", "show devtools domains"),
    ("EXPLORE", "New Categories / Niches Browse", "browse premium brandables"),
    ("EXPLORE", "New Categories / Niches Browse", "explore .io listings"),
    ("EXPLORE", "New Categories / Niches Browse", "look at gaming domains"),
    ("EXPLORE", "New Categories / Niches Browse", "browse climate tech names"),
    ("EXPLORE", "New Categories / Niches Browse", "show me edtech domains"),
    ("EXPLORE", "New Categories / Niches Browse", "browse ai domains but also show me if any have traffic or backlinks because i want value not just names"),
    ("EXPLORE", "New Categories / Niches Browse", "show devtools domains but also any saas or cloud names too because my startup could go either direction"),
    # Time-Based Browse
    ("EXPLORE", "Time-Based Browse", "what dropped today"),
    ("EXPLORE", "Time-Based Browse", "domains added in last hour"),
    ("EXPLORE", "Time-Based Browse", "listings from this morning"),
    ("EXPLORE", "Time-Based Browse", "whats new since yesterday"),
    ("EXPLORE", "Time-Based Browse", "domains posted this week"),
    ("EXPLORE", "Time-Based Browse", "recent activity in marketplace"),
    ("EXPLORE", "Time-Based Browse", "latest in the last 24 hours"),
    ("EXPLORE", "Time-Based Browse", "anything new today"),
    ("EXPLORE", "Time-Based Browse", "fresh this afternoon"),
    ("EXPLORE", "Time-Based Browse", "updates from this week"),
    ("EXPLORE", "Time-Based Browse", "what dropped today but also show me anything from yesterday that nobody bought yet"),
    ("EXPLORE", "Time-Based Browse", "whats new since yesterday but also specifically anything in ai or fintech that just got listed"),
    # Watchlist / Activity
    ("EXPLORE", "Watchlist / Activity", "show me most bid on domains"),
    ("EXPLORE", "Watchlist / Activity", "domains heating up in bids"),
    ("EXPLORE", "Watchlist / Activity", "auctions with rising activity"),
    ("EXPLORE", "Watchlist / Activity", "names people keep watching"),
    ("EXPLORE", "Watchlist / Activity", "most competitive listings now"),
    ("EXPLORE", "Watchlist / Activity", "domains with sudden interest"),
    ("EXPLORE", "Watchlist / Activity", "trending up in views"),
    ("EXPLORE", "Watchlist / Activity", "active bidding right now"),
    ("EXPLORE", "Watchlist / Activity", "domains gaining attention fast"),
    ("EXPLORE", "Watchlist / Activity", "whats picking up steam"),
    ("EXPLORE", "Watchlist / Activity", "most bid on domains right now but also show me which ones are still cheap enough to win"),
    ("EXPLORE", "Watchlist / Activity", "names people keep watching but also want to see which of those are ending soon so i can act"),
    # Messy Real User Explore
    ("EXPLORE", "Messy Real User Explore", "show me cool domains"),
    ("EXPLORE", "Messy Real User Explore", "anything interesting today"),
    ("EXPLORE", "Messy Real User Explore", "whats trending rn"),
    ("EXPLORE", "Messy Real User Explore", "what hot this week"),
    ("EXPLORE", "Messy Real User Explore", "what everybody watching now"),
    ("EXPLORE", "Messy Real User Explore", "any good ai names lately"),
    ("EXPLORE", "Messy Real User Explore", "what category getting hot"),
    ("EXPLORE", "Messy Real User Explore", "browse something interesting"),
    ("EXPLORE", "Messy Real User Explore", "show me something worth looking at"),
    ("EXPLORE", "Messy Real User Explore", "just show me whats new"),
    ("EXPLORE", "Messy Real User Explore", "surprise me with cool names"),
    ("EXPLORE", "Messy Real User Explore", "anything fun to look at today"),
    ("EXPLORE", "Messy Real User Explore", "anything interesting today but also specifically ending soon because i want to buy something today"),
    ("EXPLORE", "Messy Real User Explore", "whats trending rn but also any new listings today i want to see fresh and popular at same time"),
    # Misspelled Explore
    ("EXPLORE", "Misspelled Explore", "whats trendin right now"),
    ("EXPLORE", "Misspelled Explore", "hot doamins this week"),
    ("EXPLORE", "Misspelled Explore", "new domian added today"),
    ("EXPLORE", "Misspelled Explore", "auctons ending soon"),
    ("EXPLORE", "Misspelled Explore", "most viewd domains"),
    ("EXPLORE", "Misspelled Explore", "fresh listins worth checking"),
    ("EXPLORE", "Misspelled Explore", "anyting interesting today"),
    ("EXPLORE", "Misspelled Explore", "populer ai domains lately"),
    ("EXPLORE", "Misspelled Explore", "domians ending tonight"),
    ("EXPLORE", "Misspelled Explore", "show me cool doamins"),
    ("EXPLORE", "Misspelled Explore", "whats hapening today"),
    ("EXPLORE", "Misspelled Explore", "latst premium listings"),
    ("EXPLORE", "Misspelled Explore", "endng this weekend"),
    ("EXPLORE", "Misspelled Explore", "most actve auctions"),
    ("EXPLORE", "Misspelled Explore", "whats trendin in ai and fintec domians this week"),
    ("EXPLORE", "Misspelled Explore", "auctons endng soon under 1k worth biddin on"),
    # Incomplete Explore
    ("EXPLORE", "Incomplete Explore", "whats trending"),
    ("EXPLORE", "Incomplete Explore", "new today"),
    ("EXPLORE", "Incomplete Explore", "ending soon"),
    ("EXPLORE", "Incomplete Explore", "most viewed"),
    ("EXPLORE", "Incomplete Explore", "hot this week"),
    ("EXPLORE", "Incomplete Explore", "fresh listings"),
    ("EXPLORE", "Incomplete Explore", "anything new"),
    ("EXPLORE", "Incomplete Explore", "popular now"),
    ("EXPLORE", "Incomplete Explore", "trending in ai and"),
    ("EXPLORE", "Incomplete Explore", "ending soon under"),
    # Long Natural Explore
    ("EXPLORE", "Long Natural Explore", "can you show me whats trending in the domain market this week"),
    ("EXPLORE", "Long Natural Explore", "i just want to browse some cool startup names that got listed recently"),
    ("EXPLORE", "Long Natural Explore", "show me the auctions that are ending in the next few hours please"),
    ("EXPLORE", "Long Natural Explore", "what domains are people watching and bidding on the most right now"),
    ("EXPLORE", "Long Natural Explore", "i would like to see the freshest premium domains added today"),
    ("EXPLORE", "Long Natural Explore", "what categories are getting hot in the marketplace lately"),
    ("EXPLORE", "Long Natural Explore", "just show me something interesting to look at i dont have anything specific"),
    ("EXPLORE", "Long Natural Explore", "which ai domains are everyone talking about this week"),
    ("EXPLORE", "Long Natural Explore", "what new listings came in since yesterday worth checking"),
    ("EXPLORE", "Long Natural Explore", "show me the most active auctions happening right now"),
    ("EXPLORE", "Long Natural Explore", "i want to see what is trending right now and also if there are any new listings that came in today worth looking at"),
    ("EXPLORE", "Long Natural Explore", "show me whats hot this week and also if anything interesting dropped today that nobody noticed yet"),

    # ── GUIDANCE ─────────────────────────────────────────────────────────────
    # Startup / Founder
    ("GUIDANCE", "Startup / Founder", "starting ai startup what kind domain should i buy"),
    ("GUIDANCE", "Startup / Founder", "should startup spend money on premium domain early"),
    ("GUIDANCE", "Startup / Founder", "brandable name or exact company name better"),
    ("GUIDANCE", "Startup / Founder", "good domain strategy for b2b saas"),
    ("GUIDANCE", "Startup / Founder", "what makes startup domain memorable"),
    ("GUIDANCE", "Startup / Founder", "common founder mistakes buying domains"),
    ("GUIDANCE", "Startup / Founder", "i am starting ai startup should i get .com or .io and also does brand name or domain come first"),
    ("GUIDANCE", "Startup / Founder", "starting b2b saas should i spend 2k on good domain or is that too much for early stage startup"),
    # Small Business
    ("GUIDANCE", "Small Business", "opening local business how choose domain"),
    ("GUIDANCE", "Small Business", "should local business include city name"),
    ("GUIDANCE", "Small Business", "best domain for service business"),
    ("GUIDANCE", "Small Business", "coffee shop or restaurant domain advice"),
    ("GUIDANCE", "Small Business", "small business domain on budget"),
    ("GUIDANCE", "Small Business", "how important is .com for local business"),
    ("GUIDANCE", "Small Business", "opening local business should i use city name in domain or not and also .com or something else"),
    ("GUIDANCE", "Small Business", "small business under budget but also should i go for exact match or something more brandable"),
    # Professional / Personal Brand
    ("GUIDANCE", "Professional / Personal Brand", "consultant domain or personal name domain"),
    ("GUIDANCE", "Professional / Personal Brand", "best domain for lawyer accountant advisor"),
    ("GUIDANCE", "Professional / Personal Brand", "trust building professional domain advice"),
    ("GUIDANCE", "Professional / Personal Brand", "should i use my own name as domain"),
    ("GUIDANCE", "Professional / Personal Brand", "domain for freelancer or consultant"),
    ("GUIDANCE", "Professional / Personal Brand", "best tld for professional services"),
    ("GUIDANCE", "Professional / Personal Brand", "consultant should i use my name as domain or create brand and also which tld looks most professional"),
    ("GUIDANCE", "Professional / Personal Brand", "lawyer or accountant domain advice and also should i include practice area in name or not"),
    # Creator / Content
    ("GUIDANCE", "Creator / Content", "podcast domain recommendations"),
    ("GUIDANCE", "Creator / Content", "personal website domain advice"),
    ("GUIDANCE", "Creator / Content", "should creator buy own name domain"),
    ("GUIDANCE", "Creator / Content", "best domain for newsletter business"),
    ("GUIDANCE", "Creator / Content", "freelance portfolio domain help"),
    ("GUIDANCE", "Creator / Content", "personal brand domain strategy"),
    ("GUIDANCE", "Creator / Content", "starting podcast should i use my name or podcast title as domain and also what tld is best"),
    ("GUIDANCE", "Creator / Content", "building newsletter brand should domain match newsletter name and also what tld would you recommend"),
    # Investor
    ("GUIDANCE", "Investor", "best domain investment under 1000"),
    ("GUIDANCE", "Investor", ".com or .io better investment today"),
    ("GUIDANCE", "Investor", "are short domains still worth buying"),
    ("GUIDANCE", "Investor", "one word domains worth premium prices"),
    ("GUIDANCE", "Investor", "what makes domain valuable for resale"),
    ("GUIDANCE", "Investor", "domain investing still worth it"),
    ("GUIDANCE", "Investor", "is .com or .io better investment and also does category like ai or fintech matter for returns"),
    ("GUIDANCE", "Investor", "what makes domain valuable for resale and also is there difference between brandable and keyword value"),
    # SEO
    ("GUIDANCE", "SEO", "aged domain or fresh domain for seo"),
    ("GUIDANCE", "SEO", "does domain age matter anymore"),
    ("GUIDANCE", "SEO", "exact match or brandable domain for seo"),
    ("GUIDANCE", "SEO", "should i buy expired domain for rankings"),
    ("GUIDANCE", "SEO", "how important are backlinks when buying domain"),
    ("GUIDANCE", "SEO", "how to evaluate seo value of domain"),
    ("GUIDANCE", "SEO", "how do i check if a domains backlinks are good"),
    ("GUIDANCE", "SEO", "is it worth paying extra for a high authority domain"),
    ("GUIDANCE", "SEO", "aged domain vs fresh for seo but also does previous owner matter or just the backlink profile"),
    ("GUIDANCE", "SEO", "should i buy expired domain for rankings and also how do i know if it was ever penalized by google"),
    # Expired Domains
    ("GUIDANCE", "Expired Domains", "expired domain or auction domain better"),
    ("GUIDANCE", "Expired Domains", "when is expired domain worth buying"),
    ("GUIDANCE", "Expired Domains", "what should i check before buying expired domain"),
    ("GUIDANCE", "Expired Domains", "should i backorder this domain"),
    ("GUIDANCE", "Expired Domains", "buy in auction or wait for drop"),
    ("GUIDANCE", "Expired Domains", "red flags in expired domains"),
    ("GUIDANCE", "Expired Domains", "should i backorder or bid in auction and also what is risk difference between two approaches"),
    ("GUIDANCE", "Expired Domains", "buy expired domain for seo or brandability and also can it serve both purposes or do i have to choose"),
    # Portfolio / Advanced
    ("GUIDANCE", "Portfolio / Advanced", "how should i build domain portfolio"),
    ("GUIDANCE", "Portfolio / Advanced", "diversify across tlds or focus on .com"),
    ("GUIDANCE", "Portfolio / Advanced", "quality or quantity for domain portfolio"),
    ("GUIDANCE", "Portfolio / Advanced", "how do pros find undervalued domains"),
    ("GUIDANCE", "Portfolio / Advanced", "when should investor sell domains"),
    ("GUIDANCE", "Portfolio / Advanced", "domain portfolio mistakes to avoid"),
    ("GUIDANCE", "Portfolio / Advanced", "quality vs quantity in domain portfolio and also what is realistic roi to expect from investing"),
    ("GUIDANCE", "Portfolio / Advanced", "when to sell domains in portfolio and also how do i know if i should hold or flip specific name"),
    # Naming / Brand Strategy
    ("GUIDANCE", "Naming / Brand Strategy", "how short should domain be"),
    ("GUIDANCE", "Naming / Brand Strategy", "should domain match company name exactly"),
    ("GUIDANCE", "Naming / Brand Strategy", "are numbers or hyphens bad"),
    ("GUIDANCE", "Naming / Brand Strategy", "short name vs descriptive name"),
    ("GUIDANCE", "Naming / Brand Strategy", "should i invent a new word for brand"),
    ("GUIDANCE", "Naming / Brand Strategy", "how do i know domain sounds premium"),
    ("GUIDANCE", "Naming / Brand Strategy", "should domain match company name and also what if best .com is taken should i go .io or change name"),
    ("GUIDANCE", "Naming / Brand Strategy", "short vs descriptive domain name and also how do i know which is better for my specific business"),
    # Beginner
    ("GUIDANCE", "Beginner", "buying first domain what should i know"),
    ("GUIDANCE", "Beginner", "how much should beginner spend on first domain"),
    ("GUIDANCE", "Beginner", "buy now or auction for beginner"),
    ("GUIDANCE", "Beginner", "safest domain extension for beginner"),
    ("GUIDANCE", "Beginner", "common mistakes new buyers make"),
    ("GUIDANCE", "Beginner", "can beginners make money flipping domains"),
    ("GUIDANCE", "Beginner", "buying first domain what to know and also should i buy at auction or just go for buy now listing"),
    ("GUIDANCE", "Beginner", "common mistakes to avoid and also what are biggest mistakes experienced investors made early on"),
    # Comparison / Decision Support
    ("GUIDANCE", "Comparison / Decision Support", "which of these domains should i buy"),
    ("GUIDANCE", "Comparison / Decision Support", "is this domain worth the asking price"),
    ("GUIDANCE", "Comparison / Decision Support", "premium domain or cheaper alternative"),
    ("GUIDANCE", "Comparison / Decision Support", "would you buy this domain today"),
    ("GUIDANCE", "Comparison / Decision Support", "is this a good investment domain"),
    ("GUIDANCE", "Comparison / Decision Support", "help me choose between these domains"),
    ("GUIDANCE", "Comparison / Decision Support", "help me choose between two domains one is more expensive but shorter is it worth premium"),
    ("GUIDANCE", "Comparison / Decision Support", "comparing two domains one has traffic other is more brandable which matters more and why"),
    # Legal / Transfer / Safety
    ("GUIDANCE", "Legal / Transfer / Safety", "is it safe to buy domain from auction"),
    ("GUIDANCE", "Legal / Transfer / Safety", "how does domain transfer work"),
    ("GUIDANCE", "Legal / Transfer / Safety", "what fees come with buying a domain"),
    ("GUIDANCE", "Legal / Transfer / Safety", "can i get scammed buying expired domains"),
    ("GUIDANCE", "Legal / Transfer / Safety", "should i worry about trademark on a name"),
    ("GUIDANCE", "Legal / Transfer / Safety", "how to check if a domain was penalized before"),
    ("GUIDANCE", "Legal / Transfer / Safety", "is escrow needed for domain purchase"),
    ("GUIDANCE", "Legal / Transfer / Safety", "what happens if i miss a renewal"),
    ("GUIDANCE", "Legal / Transfer / Safety", "is it safe to buy at auction and also what is transfer process and how long does it take"),
    ("GUIDANCE", "Legal / Transfer / Safety", "should i worry about trademark on name and also how do i check if domain was previously penalized"),
    # Pricing / Negotiation
    ("GUIDANCE", "Pricing / Negotiation", "how do i know if a domain is overpriced"),
    ("GUIDANCE", "Pricing / Negotiation", "can i negotiate domain price"),
    ("GUIDANCE", "Pricing / Negotiation", "whats a fair price for a one word .com"),
    ("GUIDANCE", "Pricing / Negotiation", "how much do premium domains usually cost"),
    ("GUIDANCE", "Pricing / Negotiation", "should i lowball a domain offer"),
    ("GUIDANCE", "Pricing / Negotiation", "why are some short domains so expensive"),
    ("GUIDANCE", "Pricing / Negotiation", "how is domain value calculated"),
    ("GUIDANCE", "Pricing / Negotiation", "when is a domain price too good to be true"),
    ("GUIDANCE", "Pricing / Negotiation", "can i negotiate domain price and also how much below asking is reasonable to offer without insulting"),
    ("GUIDANCE", "Pricing / Negotiation", "when is domain price too good to be true and also what should i check when something looks suspiciously cheap"),
    # Messy Real User Queries
    ("GUIDANCE", "Messy Real User Queries", "should i buy this domain or not"),
    ("GUIDANCE", "Messy Real User Queries", "worth spending 1k on domain?"),
    ("GUIDANCE", "Messy Real User Queries", "startup domain help pls"),
    ("GUIDANCE", "Messy Real User Queries", "what domain should i get for my business"),
    ("GUIDANCE", "Messy Real User Queries", "good domain investment right now?"),
    ("GUIDANCE", "Messy Real User Queries", "better expired or auction domain"),
    ("GUIDANCE", "Messy Real User Queries", "is .io still worth buying"),
    ("GUIDANCE", "Messy Real User Queries", "premium domain overpriced or not"),
    ("GUIDANCE", "Messy Real User Queries", "how much should i pay for domain"),
    ("GUIDANCE", "Messy Real User Queries", "what would you buy if starting today"),
    ("GUIDANCE", "Messy Real User Queries", "worth buying aged domain?"),
    ("GUIDANCE", "Messy Real User Queries", "domain too long maybe?"),
    ("GUIDANCE", "Messy Real User Queries", "company name and domain same or no"),
    ("GUIDANCE", "Messy Real User Queries", "need domain for my startup not sure if should go premium or cheap and what extension matters"),
    ("GUIDANCE", "Messy Real User Queries", "worth spending like 2k on domain for startup or should that money go to marketing or product first"),
    # Misspelled Guidance
    ("GUIDANCE", "Misspelled Guidance", "shud i buy this domain or not"),
    ("GUIDANCE", "Misspelled Guidance", "is .io stil worth buying"),
    ("GUIDANCE", "Misspelled Guidance", "how much shoud i pay for domain"),
    ("GUIDANCE", "Misspelled Guidance", "startp domain help please"),
    ("GUIDANCE", "Misspelled Guidance", "wort spending 1k on a domain"),
    ("GUIDANCE", "Misspelled Guidance", "better expird or auction domain"),
    ("GUIDANCE", "Misspelled Guidance", "whats a good domian investment now"),
    ("GUIDANCE", "Misspelled Guidance", "shoud company name and domain match"),
    ("GUIDANCE", "Misspelled Guidance", "is aged domian worth buying"),
    ("GUIDANCE", "Misspelled Guidance", "how to evaluat seo value of domain"),
    ("GUIDANCE", "Misspelled Guidance", "comon mistakes new buyers make"),
    ("GUIDANCE", "Misspelled Guidance", "premiun domain overpriced or not"),
    ("GUIDANCE", "Misspelled Guidance", "shud i by this domian or not its short but no trafic is it worth 1500"),
    ("GUIDANCE", "Misspelled Guidance", "is expird domian with backlinks bettr than fresh .com for seo purposes"),
    # Long Natural Guidance
    ("GUIDANCE", "Long Natural Guidance", "i am launching a small coffee shop and not sure what domain name to pick"),
    ("GUIDANCE", "Long Natural Guidance", "should i spend a lot on a premium domain when my startup is just starting out"),
    ("GUIDANCE", "Long Natural Guidance", "i keep seeing aged domains for sale are they actually better for seo or not"),
    ("GUIDANCE", "Long Natural Guidance", "how do i tell if a domain im looking at is worth the price they are asking"),
    ("GUIDANCE", "Long Natural Guidance", "as a beginner investor where should i even start with buying domains"),
    ("GUIDANCE", "Long Natural Guidance", "is it smarter to buy a domain at auction or wait for it to drop and grab it"),
    ("GUIDANCE", "Long Natural Guidance", "i have a consultant business should i use my own name or a brandable domain"),
    ("GUIDANCE", "Long Natural Guidance", "whats the safest way to buy an expired domain without getting burned"),
    ("GUIDANCE", "Long Natural Guidance", "how do i build a domain portfolio without spending too much money upfront"),
    ("GUIDANCE", "Long Natural Guidance", "should my company name and my domain name be exactly the same thing or not"),
    ("GUIDANCE", "Long Natural Guidance", "i found a short .com that seems cheap is there usually a catch with those"),
    ("GUIDANCE", "Long Natural Guidance", "how much should i realistically budget for my very first domain purchase"),
    ("GUIDANCE", "Long Natural Guidance", "i am beginner domain investor and want to know if i should buy premium one word .com under 3k or get several cheaper names to build portfolio"),
    ("GUIDANCE", "Long Natural Guidance", "i am trying to decide between two domains one is aged .com with traffic and one is fresh short brandable .io which one would you pick and why"),

    # ── ANALYTICS ────────────────────────────────────────────────────────────
    # Marketplace Analytics
    ("ANALYTICS", "Marketplace Analytics", "how many auctions are ending this week"),
    ("ANALYTICS", "Marketplace Analytics", "total listings added last month"),
    ("ANALYTICS", "Marketplace Analytics", "average current bid price across active auctions"),
    ("ANALYTICS", "Marketplace Analytics", "new listing trend last 90 days"),
    ("ANALYTICS", "Marketplace Analytics", "auction vs buynow listing count"),
    ("ANALYTICS", "Marketplace Analytics", "percentage of auctions with at least one bid"),
    ("ANALYTICS", "Marketplace Analytics", "how many auctions ending this week and also what is total bid volume across all active listings"),
    ("ANALYTICS", "Marketplace Analytics", "new listing trend last 90 days broken down by tld and category"),
    # TLD Analytics
    ("ANALYTICS", "TLD Analytics", "top tld by listing count"),
    ("ANALYTICS", "TLD Analytics", ".com listing count and average bid last 30 days"),
    ("ANALYTICS", "TLD Analytics", "average current bid .io domains"),
    ("ANALYTICS", "TLD Analytics", ".co vs .io by listing count and average bid"),
    ("ANALYTICS", "TLD Analytics", "tld with most new listings recently"),
    ("ANALYTICS", "TLD Analytics", "tld comparison by bid activity rate"),
    ("ANALYTICS", "TLD Analytics", "how does .com compare to .io by listing count and average current bid"),
    ("ANALYTICS", "TLD Analytics", "tld with most new listings and also which category drives the most listings in that tld"),
    # Category Analytics
    ("ANALYTICS", "Category Analytics", "ai domain listing count and average bid"),
    ("ANALYTICS", "Category Analytics", "fintech domain listing activity by month"),
    ("ANALYTICS", "Category Analytics", "healthcare domain listing count recently"),
    ("ANALYTICS", "Category Analytics", "category with most active listings right now"),
    ("ANALYTICS", "Category Analytics", "category with most bids"),
    ("ANALYTICS", "Category Analytics", "category by listing count and average current bid"),
    ("ANALYTICS", "Category Analytics", "ai listing count and average bid compared to fintech and healthcare categories"),
    ("ANALYTICS", "Category Analytics", "hottest category by bid activity and which has most competitive auctions"),
    # SEO Analytics
    ("ANALYTICS", "SEO Analytics", "does traffic correlate with current bid price"),
    ("ANALYTICS", "SEO Analytics", "does higher domain authority correlate with higher current bid"),
    ("ANALYTICS", "SEO Analytics", "average domain authority across active listings"),
    ("ANALYTICS", "SEO Analytics", "zero traffic domains with current bid above 1000"),
    ("ANALYTICS", "SEO Analytics", "backlinks vs current bid price correlation"),
    ("ANALYTICS", "SEO Analytics", "average bid by domain authority tier"),
    ("ANALYTICS", "SEO Analytics", "does traffic or domain authority have bigger impact on current bid price"),
    ("ANALYTICS", "SEO Analytics", "which seo metric correlates most strongly with current bid price"),
    # Auction Analytics
    ("ANALYTICS", "Auction Analytics", "average bids per auction"),
    ("ANALYTICS", "Auction Analytics", "auctions ending with no bids"),
    ("ANALYTICS", "Auction Analytics", "bid count vs current price"),
    ("ANALYTICS", "Auction Analytics", "most competitive auctions"),
    ("ANALYTICS", "Auction Analytics", "percentage of auctions with at least one bid"),
    ("ANALYTICS", "Auction Analytics", "average current price on contested auctions"),
    ("ANALYTICS", "Auction Analytics", "average bids per auction and also what bid count typically predicts high current price"),
    ("ANALYTICS", "Auction Analytics", "auction bid rate and also does starting price affect how many bids auctions get"),
    # Expiry Analytics
    ("ANALYTICS", "Expiry Analytics", "pending delete domains expiring this week"),
    ("ANALYTICS", "Expiry Analytics", "registrar drop volume"),
    ("ANALYTICS", "Expiry Analytics", "pending delete count today"),
    ("ANALYTICS", "Expiry Analytics", "grace period domain count currently"),
    ("ANALYTICS", "Expiry Analytics", "pending delete domains with active bid interest"),
    ("ANALYTICS", "Expiry Analytics", "expiry status distribution by listing month"),
    ("ANALYTICS", "Expiry Analytics", "pending delete domains with bids and which registrars have most competitive drops"),
    ("ANALYTICS", "Expiry Analytics", "expiry status distribution and also is pending delete volume growing vs active listing volume"),
    # Investor Analytics
    ("ANALYTICS", "Investor Analytics", "average govalue to current price gap"),
    ("ANALYTICS", "Investor Analytics", "govalue premium by category"),
    ("ANALYTICS", "Investor Analytics", "govalue distribution by domain age"),
    ("ANALYTICS", "Investor Analytics", "domains where govalue score is high but current price is low"),
    ("ANALYTICS", "Investor Analytics", "category with highest average govalue score"),
    ("ANALYTICS", "Investor Analytics", "average govalue score by tld"),
    ("ANALYTICS", "Investor Analytics", "govalue premium by category and which one has the best value gaps"),
    ("ANALYTICS", "Investor Analytics", "category with highest govalue and also is that consistent across all auction types"),
    # Buynow / Entry Analytics
    ("ANALYTICS", "Buynow / Entry Analytics", "buynow and closeout listing count"),
    ("ANALYTICS", "Buynow / Entry Analytics", "average price on buynow listings"),
    ("ANALYTICS", "Buynow / Entry Analytics", "low price domain count under 100"),
    ("ANALYTICS", "Buynow / Entry Analytics", "single bid auctions ending soon"),
    ("ANALYTICS", "Buynow / Entry Analytics", "tld distribution on buynow listings"),
    ("ANALYTICS", "Buynow / Entry Analytics", "new listings added per week trend"),
    ("ANALYTICS", "Buynow / Entry Analytics", "buynow listing count by tld and which tld has most low price listings"),
    ("ANALYTICS", "Buynow / Entry Analytics", "new listings per week trend and also which category gets most new listings recently"),
    # Professional Category Analytics
    ("ANALYTICS", "Professional Category Analytics", "professional category domain listing count"),
    ("ANALYTICS", "Professional Category Analytics", "average bid for healthcare and legal category domains"),
    ("ANALYTICS", "Professional Category Analytics", "healthcare category listings by month"),
    ("ANALYTICS", "Professional Category Analytics", "healthcare category domain listing trend"),
    ("ANALYTICS", "Professional Category Analytics", "tld distribution for healthcare and legal category domains"),
    ("ANALYTICS", "Professional Category Analytics", "high domain authority domains in healthcare and legal categories"),
    ("ANALYTICS", "Professional Category Analytics", "healthcare category listing trend by month and also is demand growing or flat"),
    ("ANALYTICS", "Professional Category Analytics", "tld distribution for professional categories and also how does .com compare to .io"),
    # High Value Domain Analytics
    ("ANALYTICS", "High Value Domain Analytics", "high da pending delete domain count"),
    ("ANALYTICS", "High Value Domain Analytics", "auctions with 5 or more bids"),
    ("ANALYTICS", "High Value Domain Analytics", "high domain authority listings in pending delete"),
    ("ANALYTICS", "High Value Domain Analytics", "domains with high govalue and high domain authority"),
    ("ANALYTICS", "High Value Domain Analytics", "domains where govalue is high but current price is low"),
    ("ANALYTICS", "High Value Domain Analytics", "pending delete domain distribution by registrar"),
    ("ANALYTICS", "High Value Domain Analytics", "high competition auctions with 5 or more bids and also average price premium vs single bid auctions"),
    ("ANALYTICS", "High Value Domain Analytics", "high value pending delete domains by domain authority and which registrar has the most"),
    # Rankings / Leaderboards
    ("ANALYTICS", "Rankings / Leaderboards", "top current bids this month"),
    ("ANALYTICS", "Rankings / Leaderboards", "top categories by total bid volume"),
    ("ANALYTICS", "Rankings / Leaderboards", "top tlds by listing volume"),
    ("ANALYTICS", "Rankings / Leaderboards", "top domains by traffic in active listings"),
    ("ANALYTICS", "Rankings / Leaderboards", "top domains by govalue to price ratio"),
    ("ANALYTICS", "Rankings / Leaderboards", "top categories by average govalue score"),
    ("ANALYTICS", "Rankings / Leaderboards", "top categories by listing count and also which has highest average current bid"),
    ("ANALYTICS", "Rankings / Leaderboards", "top tlds by listing volume and also which tld has highest average current bid"),
    # Price / Distribution Analytics
    ("ANALYTICS", "Price / Distribution Analytics", "price distribution of active auction listings"),
    ("ANALYTICS", "Price / Distribution Analytics", "median current bid this month"),
    ("ANALYTICS", "Price / Distribution Analytics", "how many active listings have current bid above 5k"),
    ("ANALYTICS", "Price / Distribution Analytics", "percentage of listings under 500"),
    ("ANALYTICS", "Price / Distribution Analytics", "average price by domain length"),
    ("ANALYTICS", "Price / Distribution Analytics", "one word vs two word price gap"),
    ("ANALYTICS", "Price / Distribution Analytics", "share of active listings over 10k"),
    ("ANALYTICS", "Price / Distribution Analytics", "price range with most active listings"),
    ("ANALYTICS", "Price / Distribution Analytics", "median vs average current bid this month and also what does the spread tell us"),
    ("ANALYTICS", "Price / Distribution Analytics", "average bid by domain length and also are there exceptions where longer domains bid higher"),
    # Time / Trend Analytics
    ("ANALYTICS", "Time / Trend Analytics", "auctions ending by day of week"),
    ("ANALYTICS", "Time / Trend Analytics", "month with most auctions ending this year"),
    ("ANALYTICS", "Time / Trend Analytics", "month over month new listings growth"),
    ("ANALYTICS", "Time / Trend Analytics", "listing volume this year vs last year"),
    ("ANALYTICS", "Time / Trend Analytics", "seasonal trend in auction listings by month"),
    ("ANALYTICS", "Time / Trend Analytics", "listing volume by quarter"),
    ("ANALYTICS", "Time / Trend Analytics", "auctions ending weekend vs weekday"),
    ("ANALYTICS", "Time / Trend Analytics", "average days until auction ends from listing date"),
    ("ANALYTICS", "Time / Trend Analytics", "month over month new listings growth and also is growth in listing count or higher average bids"),
    ("ANALYTICS", "Time / Trend Analytics", "seasonal listing trend and also what time of year has most auctions ending"),
    # Length / Keyword Analytics
    ("ANALYTICS", "Length / Keyword Analytics", "average domain name length in active listings"),
    ("ANALYTICS", "Length / Keyword Analytics", "do shorter domains get more bids"),
    ("ANALYTICS", "Length / Keyword Analytics", "keyword frequency in high bid listings"),
    ("ANALYTICS", "Length / Keyword Analytics", "most common words in active listings"),
    ("ANALYTICS", "Length / Keyword Analytics", "do domains with numbers have lower current bids"),
    ("ANALYTICS", "Length / Keyword Analytics", "hyphen vs non-hyphen average current bid"),
    ("ANALYTICS", "Length / Keyword Analytics", "four letter domain average current bid"),
    ("ANALYTICS", "Length / Keyword Analytics", "single dictionary word domain bid premium"),
    ("ANALYTICS", "Length / Keyword Analytics", "do shorter domains get more bids and also do they bid higher or just attract more bidders"),
    ("ANALYTICS", "Length / Keyword Analytics", "do domains with numbers have lower bids and also does hyphen impact bid price similarly"),
    # Messy Real User Analytics Queries
    ("ANALYTICS", "Messy Real User Analytics Queries", "whats avg current bid across active listings"),
    ("ANALYTICS", "Messy Real User Analytics Queries", "top tld by listing count last month"),
    ("ANALYTICS", "Messy Real User Analytics Queries", "how many ai domains listed right now"),
    ("ANALYTICS", "Messy Real User Analytics Queries", "show new listing trend last 90 days"),
    ("ANALYTICS", "Messy Real User Analytics Queries", "any growth in fintech domain listings"),
    ("ANALYTICS", "Messy Real User Analytics Queries", "does traffic help bid price"),
    ("ANALYTICS", "Messy Real User Analytics Queries", "most bid auctions this week"),
    ("ANALYTICS", "Messy Real User Analytics Queries", "how many pending delete domains right now"),
    ("ANALYTICS", "Messy Real User Analytics Queries", "which categories have highest govalue avg"),
    ("ANALYTICS", "Messy Real User Analytics Queries", "are buynow listing counts growing"),
    ("ANALYTICS", "Messy Real User Analytics Queries", "which domains have highest govalue vs price gap"),
    ("ANALYTICS", "Messy Real User Analytics Queries", "top category by total bid volume"),
    ("ANALYTICS", "Messy Real User Analytics Queries", "show me marketplace stats summary"),
    ("ANALYTICS", "Messy Real User Analytics Queries", "whats avg current bid lately and is it going up or down and which categories are driving that"),
    ("ANALYTICS", "Messy Real User Analytics Queries", "does traffic help bid price and also is da or backlinks bigger factor than raw traffic"),
    # Misspelled Analytics
    ("ANALYTICS", "Misspelled Analytics", "how many auctons ending this week"),
    ("ANALYTICS", "Misspelled Analytics", "whats avg curnt bid pric lately"),
    ("ANALYTICS", "Misspelled Analytics", "top tld by listng cnt last month"),
    ("ANALYTICS", "Misspelled Analytics", "how many ai domians listed now"),
    ("ANALYTICS", "Misspelled Analytics", "show new listngs trnd last 90 days"),
    ("ANALYTICS", "Misspelled Analytics", "any growth in fintec domain listngs"),
    ("ANALYTICS", "Misspelled Analytics", "does trafic help bid pric"),
    ("ANALYTICS", "Misspelled Analytics", "does higher da sel for more"),
    ("ANALYTICS", "Misspelled Analytics", "most bidded auctons this week"),
    ("ANALYTICS", "Misspelled Analytics", "how many expird or pendig delet domains now"),
    ("ANALYTICS", "Misspelled Analytics", "which categorys hav highest govalue scor"),
    ("ANALYTICS", "Misspelled Analytics", "averge govalue to pric gap"),
    ("ANALYTICS", "Misspelled Analytics", "whats avg bid pric and is it going up or down in lst 90 days"),
    ("ANALYTICS", "Misspelled Analytics", "does trafic or da hav bigger impct on curnt bid pric for domians"),
    # Long Natural Analytics
    ("ANALYTICS", "Long Natural Analytics", "can you tell me how many auctions are closing in the marketplace this week"),
    ("ANALYTICS", "Long Natural Analytics", "what has the average current bid been across active listings over the last month"),
    ("ANALYTICS", "Long Natural Analytics", "which tld has the most listings and highest total bid volume in the last thirty days"),
    ("ANALYTICS", "Long Natural Analytics", "is there any noticeable growth in ai domain listings compared to a month ago"),
    ("ANALYTICS", "Long Natural Analytics", "do domains with more traffic actually have higher current bid prices"),
    ("ANALYTICS", "Long Natural Analytics", "show me how the overall new listing trend has moved over the last ninety days"),
    ("ANALYTICS", "Long Natural Analytics", "which category has the most active listings and highest total bid volume right now"),
    ("ANALYTICS", "Long Natural Analytics", "has the number of buynow listings been growing lately or not"),
    ("ANALYTICS", "Long Natural Analytics", "what is average current bid over last month and is it trending up or down compared to three months ago"),
    ("ANALYTICS", "Long Natural Analytics", "do domains with more traffic have higher current bids and also is da or backlinks stronger predictor of bid price"),
]

EVAL_KS: List[int] = [5]


# ── IR helpers ────────────────────────────────────────────────────────────────

def _ndcg_at(scores: List[float], k: int) -> float:
    """NDCG@k — same formula as ``app._compute_retrieval_metrics`` (linear DCG).

    Uses ``coherence_score`` as the label-free relevance proxy (no gold labels).
    Items already sorted by score yield NDCG≈1.0; that matches the API metric.
    """
    top = scores[: min(k, len(scores))]
    if not top:
        return 0.0
    ranks = list(range(1, len(top) + 1))
    dcg = sum(s / math.log2(r + 1) for s, r in zip(top, ranks))
    idcg = sum(
        s / math.log2(r + 1)
        for s, r in zip(sorted(top, reverse=True), ranks)
    )
    return round(dcg / idcg, 4) if idcg > 0 else 0.0


def _prf_at(scores: List[float], k: int, threshold: float) -> Tuple[float, float, float]:
    """Precision/Recall/F1@k — same relevance rule as ``app._compute_retrieval_metrics``.

    Score >= ``threshold`` (``retrieval.metrics.relevance_threshold``) counts as
    relevant. No gold labels; pool = returned ranked list.
    """
    top = scores[: min(k, len(scores))]
    if not top:
        return 0.0, 0.0, 0.0
    rel_in_k = sum(1 for s in top if s >= threshold)
    relevant_in_pool = sum(1 for s in scores if s >= threshold)
    precision_denom = min(k, len(top))
    precision = (
        round(rel_in_k / float(precision_denom), 4) if precision_denom > 0 else 0.0
    )
    recall = (
        round(rel_in_k / float(relevant_in_pool), 4) if relevant_in_pool > 0 else 0.0
    )
    f1 = (
        round(2 * precision * recall / (precision + recall), 4)
        if (precision + recall) > 0
        else 0.0
    )
    return precision, recall, f1


# ── Data types ────────────────────────────────────────────────────────────────

@dataclass
class RankedStats:
    count: int
    zero_results: bool
    avg_coherence: Optional[float]
    min_coherence: Optional[float]
    max_coherence: Optional[float]
    top_domains: List[str]
    score_std: Optional[float]
    score_gap_ratio: Optional[float]   # top1 / mean — how much the top result stands out
    score_monotonic_pct: float         # % of consecutive pairs where score[i] >= score[i+1]


@dataclass
class QueryResult:
    api_label: str
    suite: str
    section: str
    query: str
    latency_ms: float
    started_at: float            # monotonic timestamp of when the HTTP call was issued
    answer_mode: str
    mode_correct: bool
    classified_intent: Optional[str]
    intent_confidence: Optional[float]
    decision_tier: Optional[str]
    multi_intent: bool
    hard_filter_count: int
    filter_names: List[str]
    total_candidates: int
    ndcg: Optional[float]
    coherence: Optional[float]
    recall: Optional[float]
    precision: Optional[float]
    api_metrics: Dict[str, float]          # all @K keys from retrieval_metrics
    local_ndcg_at: Dict[int, float]        # computed locally from ranked_results
    local_prec_at: Dict[int, float]
    local_recall_at: Dict[int, float]
    local_f1_at: Dict[int, float]
    local_coh_at: Dict[int, float]           # mean coherence of top-K results
    ranked: RankedStats
    timeout: bool
    sla_breach: bool
    failure_mode: Optional[str]
    backends_active: List[str]
    error: Optional[str]
    # Count of identified filters with chip_kind=="hard" (soft chips excluded).
    entities_extracted_count: int
    type_id_entities: List[str]          # auction_type entity values extracted (for precision check)
    # Copied from response query_intelligence.decision_cost_usd (QI/LLM decision path).
    decision_cost_usd: float
    # Server-reported body["latency_ms"] when present; None on client error/timeout.
    # sla_breach / ≤4s% / ≤10s% use client wall latency_ms (HTTP round-trip), not this.
    server_latency_ms: Optional[float]
    # pipeline_trace.stages when measurement.ranking_stage_attribution.include_in_response
    # is on; empty dict otherwise (capture is best-effort — never fails the query).
    pipeline_stages: Dict[str, Any]
    stage_fusion_ms: Optional[float]
    stage_eranker_ms: Optional[float]
    retrieve_sources: List[str]


def _suite_sla_ms(suite: str, fallback_ms: int) -> int:
    return SUITE_SLA_MS.get(suite, fallback_ms)


# Keyword_* filter slots whose values are lexical terms (matched with keywords arrays).
_KEYWORD_VALUE_PARAMS: frozenset = frozenset({
    "keyword_contains",
    "keyword_starts_with",
    "keyword_ends_with",
    "keyword_phrase",
    "keyword_contains_exclude",
})


def _filter_slot_name(entry: dict) -> str:
    """Resolve FIND api_param / slot name from an identified or applied_filters chip."""
    api = entry.get("api_param")
    if isinstance(api, dict) and api.get("name"):
        return str(api["name"]).strip()
    if isinstance(api, str) and api.strip():
        return api.strip()
    return str(entry.get("name") or entry.get("param") or "").strip()


def _eval_hard_filters(
    identified: List[dict],
    applied_filters: List[dict],
) -> List[dict]:
    """Hard chips for eval: prefer applied_filters; else slim identified (name+value).

    Public ``filters.identified`` no longer carries ``chip_kind`` / ``data_status``.
    ``pipeline_trace.applied_filters`` is the grounded hard set retrieval uses.
    When applied is empty (degraded / empty inventory), fall back to identified.
    ``keyword_*`` lexical chips are excluded (counted via ``_eval_keyword_terms``).
    """
    src = applied_filters if applied_filters else identified
    out: List[dict] = []
    for f in src:
        if not isinstance(f, dict):
            continue
        name = _filter_slot_name(f)
        if not name or name in _KEYWORD_VALUE_PARAMS or name == "keyword_match_mode":
            continue
        if not applied_filters:
            # Legacy identified rows may still stamp chip_kind / data_status.
            kind = f.get("chip_kind")
            if kind is not None and kind != "hard":
                continue
            if f.get("data_status") not in (None, "available"):
                continue
        out.append(f)
    return out


def _eval_keyword_terms(
    keywords: List[dict],
    applied_keywords: List[dict],
    soft_signals: List[dict],
    identified: List[dict],
    applied_filters: List[dict],
) -> List[str]:
    """Union keyword terms from keywords / applied_keywords / keyword_* filter chips."""
    terms: Set[str] = set()
    for group in (keywords, applied_keywords):
        for kw in group or []:
            if not isinstance(kw, dict):
                continue
            t = str(kw.get("term") or "").strip().lower()
            if t:
                terms.add(t)
    for entry in list(soft_signals or []) + list(identified or []) + list(applied_filters or []):
        if not isinstance(entry, dict):
            continue
        name = _filter_slot_name(entry)
        if name not in _KEYWORD_VALUE_PARAMS:
            continue
        raw = entry.get("value")
        values = raw if isinstance(raw, list) else [raw]
        for v in values:
            if v is None:
                continue
            for part in str(v).replace("|", ",").split(","):
                t = part.strip().lower()
                if t:
                    terms.add(t)
    return sorted(terms)


def _extract_pipeline_stages(
    body: Dict[str, Any],
) -> Tuple[Dict[str, Any], Optional[float], Optional[float], List[str]]:
    """Pull ranking-stage attribution from `/search` `pipeline_trace.stages` if present."""
    pt = body.get("pipeline_trace") or {}
    stages = pt.get("stages") if isinstance(pt, dict) else None
    if not isinstance(stages, dict):
        return {}, None, None, []
    fusion = stages.get("fusion_latency_ms")
    eranker = stages.get("eranker_latency_ms")
    try:
        fusion_f = float(fusion) if fusion is not None else None
    except (TypeError, ValueError):
        fusion_f = None
    try:
        eranker_f = float(eranker) if eranker is not None else None
    except (TypeError, ValueError):
        eranker_f = None
    sources_raw = stages.get("retrieve_sources") or []
    sources = [str(s) for s in sources_raw] if isinstance(sources_raw, list) else []
    return dict(stages), fusion_f, eranker_f, sources


def _empty_pipeline_fields() -> Tuple[Dict[str, Any], Optional[float], Optional[float], List[str]]:
    return {}, None, None, []


def _is_fail_row(r: QueryResult) -> bool:
    """Fail = error, client/server timeout, suite-SLA breach, or wrong answer_mode."""
    return bool(r.error or r.timeout or r.sla_breach or (not r.mode_correct))


# ── Ranked results interpretation ─────────────────────────────────────────────

def _interpret_ranked(ranked_results: List[dict]) -> RankedStats:
    if not ranked_results:
        return RankedStats(0, True, None, None, None, [], None, None, 0.0)
    scores = [r.get("coherence_score") for r in ranked_results if r.get("coherence_score") is not None]
    domains = [r.get("domain_name", "") for r in ranked_results]
    score_std: Optional[float] = round(statistics.stdev(scores), 4) if len(scores) >= 2 else None
    mean_s = statistics.mean(scores) if scores else None
    score_gap_ratio: Optional[float] = round(scores[0] / mean_s, 4) if (scores and mean_s and mean_s > 0) else None
    if len(scores) >= 2:
        pairs = [(scores[i], scores[i + 1]) for i in range(len(scores) - 1)]
        mono_count = sum(1 for a, b in pairs if a >= b)
        score_monotonic_pct = round(mono_count / len(pairs), 4)
    else:
        score_monotonic_pct = 1.0
    return RankedStats(
        count=len(ranked_results),
        zero_results=False,
        avg_coherence=statistics.mean(scores) if scores else None,
        min_coherence=min(scores) if scores else None,
        max_coherence=max(scores) if scores else None,
        top_domains=domains[:5],
        score_std=score_std,
        score_gap_ratio=score_gap_ratio,
        score_monotonic_pct=score_monotonic_pct,
    )


# ── Core search call ──────────────────────────────────────────────────────────

_EMPTY_AT: Dict[int, float] = {k: 0.0 for k in EVAL_KS}


async def search_one(
    client: httpx.AsyncClient,
    api_label: str,
    base_url: str,
    suite: str,
    section: str,
    query: str,
    sem: asyncio.Semaphore,
    top_k: int,
    sla_ms: int,
    timeout_s: float,
    rel_threshold: float = _REL_THRESHOLD,
) -> QueryResult:
    search_url = f"{base_url.rstrip('/')}/search"
    async with sem:
        t0 = time.monotonic()
        session_id = f"eval-{uuid.uuid4().hex[:12]}"
        try:
            resp = await client.post(
                search_url,
                data={"query": query, "top_k": max(top_k, max(EVAL_KS))},
                headers={"X-Session-Id": session_id},
                timeout=timeout_s,
            )
            latency_ms = round((time.monotonic() - t0) * 1000, 1)
            resp.raise_for_status()
            body: Dict[str, Any] = resp.json()
        except httpx.TimeoutException as e:
            latency_ms = round((time.monotonic() - t0) * 1000, 1)
            _ps, _sf, _se, _rs = _empty_pipeline_fields()
            return QueryResult(
                api_label=api_label, suite=suite, section=section, query=query,
                latency_ms=latency_ms, answer_mode="client_timeout", mode_correct=False,
                classified_intent=None, intent_confidence=None, decision_tier=None,
                multi_intent=False,
                hard_filter_count=0, filter_names=[], total_candidates=0,
                ndcg=None, coherence=None, recall=None, precision=None,
                api_metrics={}, local_ndcg_at=dict(_EMPTY_AT),
                local_prec_at=dict(_EMPTY_AT), local_recall_at=dict(_EMPTY_AT), local_f1_at=dict(_EMPTY_AT),
                local_coh_at=dict(_EMPTY_AT), started_at=t0,
                ranked=RankedStats(0, True, None, None, None, [], None, None, 0.0),
                timeout=True, sla_breach=True, failure_mode="client_timeout",
                backends_active=[], error=f"client_timeout: {e}",
                entities_extracted_count=0, type_id_entities=[],
                decision_cost_usd=0.0, server_latency_ms=None,
                pipeline_stages=_ps, stage_fusion_ms=_sf, stage_eranker_ms=_se,
                retrieve_sources=_rs,
            )
        except httpx.ConnectError:
            raise
        except Exception as e:
            latency_ms = round((time.monotonic() - t0) * 1000, 1)
            _ps, _sf, _se, _rs = _empty_pipeline_fields()
            return QueryResult(
                api_label=api_label, suite=suite, section=section, query=query,
                latency_ms=latency_ms, answer_mode="error", mode_correct=False,
                classified_intent=None, intent_confidence=None, decision_tier=None,
                multi_intent=False,
                hard_filter_count=0, filter_names=[], total_candidates=0,
                ndcg=None, coherence=None, recall=None, precision=None,
                api_metrics={}, local_ndcg_at=dict(_EMPTY_AT),
                local_prec_at=dict(_EMPTY_AT), local_recall_at=dict(_EMPTY_AT), local_f1_at=dict(_EMPTY_AT),
                local_coh_at=dict(_EMPTY_AT), started_at=t0,
                ranked=RankedStats(0, True, None, None, None, [], None, None, 0.0),
                timeout=False,
                sla_breach=latency_ms > _suite_sla_ms(suite, sla_ms),
                failure_mode=None,
                backends_active=[], error=str(e),
                entities_extracted_count=0, type_id_entities=[],
                decision_cost_usd=0.0, server_latency_ms=None,
                pipeline_stages=_ps, stage_fusion_ms=_sf, stage_eranker_ms=_se,
                retrieve_sources=_rs,
            )

    qi = body.get("query_intelligence") or {}
    filters_block = qi.get("filters") or {}
    pipeline_trace = body.get("pipeline_trace") or {}
    # Public identified is slim name+value (no chip_kind). applied_filters = grounded
    # hard chips retrieval uses. Union both so match covers chips on either side.
    identified = list(filters_block.get("identified") or [])
    applied_filters = list(pipeline_trace.get("applied_filters") or [])
    soft_signals = list(filters_block.get("soft_signals") or [])
    hard_filters = _eval_hard_filters(identified, applied_filters)
    kw_terms = _eval_keyword_terms(
        list(filters_block.get("keywords") or []),
        list(pipeline_trace.get("applied_keywords") or []),
        soft_signals,
        identified,
        applied_filters,
    )
    # Entity coverage = hard chips; keyword-only queries still count (encode/soft boost).
    _entities_extracted_count: int = (
        len(hard_filters) if hard_filters else len(kw_terms)
    )
    # Auction-type hard chips use FIND names typeIncludeList / typeExcludeList
    # (slot "auction_type" may appear only as legacy name before api_param resolve).
    _TYPE_PARAM_NAMES = frozenset({
        "typeIncludeList", "typeExcludeList", "auction_type",
    })
    _type_id_entities: List[str] = [
        str(e.get("value", ""))
        for e in hard_filters
        if _filter_slot_name(e) in _TYPE_PARAM_NAMES
    ]
    try:
        _decision_cost_usd: float = float(qi.get("decision_cost_usd") or 0.0)
    except (TypeError, ValueError):
        _decision_cost_usd = 0.0

    rm = body.get("retrieval_metrics") or {}
    api_metrics: Dict[str, float] = {k: float(v) for k, v in rm.items() if "@" in str(k) and isinstance(v, (int, float))}
    _eval_k  = EVAL_KS[0]
    ndcg     = api_metrics.get(f"NDCG@{_eval_k}")
    coh      = api_metrics.get(f"Coherence@{_eval_k}")
    recall_v = api_metrics.get(f"Recall@{_eval_k}")
    prec_v   = api_metrics.get(f"Precision@{_eval_k}")

    answer_mode = body.get("answer_mode", "")
    # classified_intent is needed here (before the return block) to determine
    # routing correctness for ANALYTICS independently of SQL execution outcome.
    _classified_intent_early: Optional[str] = qi.get("classified_intent")

    # For ANALYTICS, routing correctness is determined by intent classification,
    # not by whether the SQL execution layer succeeded.  A query classified as
    # 'analytics' that falls back to explore_fallback (e.g. no substrate) still
    # represents a correct routing decision — the infrastructure failure is
    # tracked separately via the sql_analytics_failed backend signal.
    if suite == "ANALYTICS":
        mode_correct = (
            answer_mode in EXPECTED_MODE.get(suite, set())
            or _classified_intent_early == "analytics"
        )
    else:
        mode_correct = answer_mode in EXPECTED_MODE.get(suite, set())

    failure_mode_val = rm.get("failure_mode")
    # A timeout is a genuine latency/deadline event only. answer_mode ==
    # 'explore_fallback' is the zero-result guard degrading gracefully (fast,
    # not a deadline breach) and is already tracked in the Fallback column —
    # counting it as a timeout conflated two unrelated outcomes for HYBRID.
    if suite == "ANALYTICS":
        is_timeout = failure_mode_val in ("timeout", "analytics_timeout", "explore_fallback_timeout")
    else:
        is_timeout = failure_mode_val == "timeout"

    ranked_results = body.get("ranked_results") or []
    r_scores = [r.get("coherence_score") for r in ranked_results if r.get("coherence_score") is not None]

    local_ndcg_at: Dict[int, float] = {}
    local_prec_at: Dict[int, float] = {}
    local_recall_at: Dict[int, float] = {}
    local_f1_at: Dict[int, float] = {}
    for k_val in EVAL_KS:
        local_ndcg_at[k_val] = _ndcg_at(r_scores, k_val)
        p, rc, f = _prf_at(r_scores, k_val, rel_threshold)
        local_prec_at[k_val]   = p
        local_recall_at[k_val] = rc
        local_f1_at[k_val]     = f

    local_coh_at: Dict[int, float] = {
        k: (statistics.mean(r_scores[:k]) if r_scores[:k] else 0.0)
        for k in EVAL_KS
    }

    # backends_active is built from ranked_results contributing_sources only.
    # For ANALYTICS the SQL pipeline runs in a separate code path and never
    # appears in contributing_sources — infer it from the analytics response block.
    _analytics_blk = body.get("analytics") or {}
    _backends_active: List[str] = list(rm.get("backends_active") or [])
    if answer_mode == "analytics":
        if _analytics_blk.get("execution"):
            _backends_active.append("sql_analytics")       # SQL ran and succeeded
        elif _analytics_blk:
            _backends_active.append("sql_analytics_failed")  # SQL attempted, failed
    elif (
        suite == "ANALYTICS"
        and _classified_intent_early == "analytics"
        and answer_mode == "explore_fallback"
    ):
        # Correctly classified as analytics but SQL substrate unavailable —
        # routing succeeded, execution layer did not.
        _backends_active.append("sql_analytics_failed")

    _pipeline_stages, _stage_fusion_ms, _stage_eranker_ms, _retrieve_sources = (
        _extract_pipeline_stages(body)
    )
    _suite_sla = _suite_sla_ms(suite, sla_ms)
    try:
        _server_latency_ms: Optional[float] = (
            float(body["latency_ms"]) if body.get("latency_ms") is not None else None
        )
    except (TypeError, ValueError):
        _server_latency_ms = None

    return QueryResult(
        api_label=api_label,
        suite=suite,
        section=section,
        query=query,
        latency_ms=latency_ms,
        answer_mode=answer_mode,
        mode_correct=mode_correct,
        classified_intent=qi.get("classified_intent"),
        intent_confidence=qi.get("intent_confidence"),
        decision_tier=qi.get("decision_tier"),
        multi_intent=bool(qi.get("multi_intent", False)),
        hard_filter_count=len(hard_filters),
        filter_names=[_filter_slot_name(f) for f in hard_filters],
        total_candidates=rm.get("total_candidates", 0),
        ndcg=float(ndcg) if ndcg is not None else None,
        coherence=float(coh) if coh is not None else None,
        recall=float(recall_v) if recall_v is not None else None,
        precision=float(prec_v) if prec_v is not None else None,
        api_metrics=api_metrics,
        local_ndcg_at=local_ndcg_at,
        local_prec_at=local_prec_at,
        local_recall_at=local_recall_at,
        local_f1_at=local_f1_at,
        local_coh_at=local_coh_at,
        started_at=t0,
        ranked=_interpret_ranked(ranked_results),
        timeout=is_timeout,
        sla_breach=latency_ms > _suite_sla,
        failure_mode=failure_mode_val,
        backends_active=_backends_active,
        error=None,
        entities_extracted_count=_entities_extracted_count,
        type_id_entities=_type_id_entities,
        decision_cost_usd=_decision_cost_usd,
        server_latency_ms=_server_latency_ms,
        pipeline_stages=_pipeline_stages,
        stage_fusion_ms=_stage_fusion_ms,
        stage_eranker_ms=_stage_eranker_ms,
        retrieve_sources=_retrieve_sources,
    )


# ── Stats helpers ─────────────────────────────────────────────────────────────

def pct(n: int, total: int) -> str:
    return f"{round(100 * n / total)}%" if total else "N/A"


def _grade(val_pct: int, warn: int, good: int, direction: str = "hi", note: str = "") -> str:
    """Return OK/WARN/BAD with benchmark label. direction='hi'=higher better, 'lo'=lower better."""
    if direction == "hi":
        sym = "OK" if val_pct >= good else ("WARN" if val_pct >= warn else "BAD")
        label = f"target >={good}%"
    else:
        sym = "OK" if val_pct <= good else ("WARN" if val_pct <= warn else "BAD")
        label = f"target <={good}%"
    return f"{sym} [{label}{'; ' + note if note else ''}]"


def _lat_p95_grade(p95_ms: float, warn_ms: float, good_ms: float, note: str = "") -> str:
    """Latency grade (lower better): OK if p95<=good, WARN if p95<=warn, else BAD."""
    if p95_ms <= good_ms:
        sym = "OK"
    elif p95_ms <= warn_ms:
        sym = "WARN"
    else:
        sym = "BAD"
    extra = f"; {note}" if note else ""
    return f"{sym} [p95<={good_ms:.0f}ms target; WARN<={warn_ms:.0f}ms{extra}]"


def lat_stats(vals: List[float]) -> str:
    if not vals:
        return "N/A"
    sv = sorted(vals)
    n = len(sv)
    return (
        f"p50={sv[n // 2]:.0f}ms  "
        f"p90={sv[min(int(n * 0.90), n - 1)]:.0f}ms  "
        f"p95={sv[min(int(n * 0.95), n - 1)]:.0f}ms  "
        f"p99={sv[min(int(n * 0.99), n - 1)]:.0f}ms  "
        f"mean={statistics.mean(vals):.0f}ms  "
        f"max={max(vals):.0f}ms"
    )


def _mms(vals: List[float], fmt: str = ".4f") -> str:
    if not vals:
        return "N/A"
    return f"mean={statistics.mean(vals):{fmt}}  min={min(vals):{fmt}}  max={max(vals):{fmt}}"


def _prf_row(qs: List[QueryResult], k: int) -> str:
    if not qs:
        return "P=N/A  R=N/A  F1=N/A"
    ps = [q.local_prec_at.get(k, 0.0) for q in qs]
    rs = [q.local_recall_at.get(k, 0.0) for q in qs]
    fs = [q.local_f1_at.get(k, 0.0) for q in qs]
    return (
        f"P={statistics.mean(ps):.4f}  "
        f"R={statistics.mean(rs):.4f}  "
        f"F1={statistics.mean(fs):.4f}"
    )




# ── Orchestration ─────────────────────────────────────────────────────────────

async def run_suite_queries(
    api_label: str,
    base_url: str,
    queries: List[tuple],
    concurrency: int,
    top_k: int,
    sla_ms: int,
    timeout_s: float,
    rel_threshold: float = _REL_THRESHOLD,
) -> List[QueryResult]:
    sem = asyncio.Semaphore(concurrency)
    results: List[QueryResult] = []
    async with httpx.AsyncClient() as client:
        tasks = [
            asyncio.create_task(
                search_one(client, api_label, base_url, suite, section, query, sem, top_k, sla_ms, timeout_s, rel_threshold)
            )
            for suite, section, query in queries
        ]
        done = 0
        t_start = time.perf_counter()
        for fut in asyncio.as_completed(tasks):
            try:
                r = await fut
            except httpx.ConnectError as e:
                for t in tasks:
                    t.cancel()
                sys.stderr.write(f"\n  [{api_label}] Server unreachable ({e}). Aborting.\n")
                sys.stderr.flush()
                raise
            results.append(r)
            done += 1
            if done % 50 == 0 or done == len(tasks):
                elapsed = time.perf_counter() - t_start
                qps = done / elapsed if elapsed > 0 else 0.0
                remaining = (len(tasks) - done) / qps if qps > 0 else 0.0
                sys.stderr.write(
                    f"  [{api_label}] {done}/{len(tasks)} queries"
                    f"  elapsed={elapsed:.0f}s  qps={qps:.1f}"
                    f"  eta={remaining:.0f}s\n"
                )
                sys.stderr.flush()
    order = {(s, sec, q): i for i, (s, sec, q) in enumerate(queries)}
    results.sort(key=lambda r: order.get((r.suite, r.section, r.query), 9999))
    return results



def _write_markdown_report(
    results: List[QueryResult],
    api: str,
    wall_s: float,
    peak_mb: float,
    total_queries_run: int,
    sla_ms: int,
    out_dir: Path = DEFAULT_OUTPUT_DIR,
    stamp: Optional[str] = None,
) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = stamp or str(int(time.time()))
    out_path = out_dir / f"eval_report_{stamp}.md"
    lines: List[str] = []

    def w(s: str = "") -> None:
        lines.append(s)

    def _agg(qs: List[QueryResult], suite: str) -> dict:
        # Shared with xlsx suite/section sheets — keeps @5 + latency + cost + guards.
        return _agg_for_report(qs, suite, sla_ms)

    def _render_api(results: List[QueryResult]) -> None:
        actual_ws = (
            max(r.started_at + r.latency_ms / 1000 for r in results)
            - min(r.started_at for r in results)
        ) if results else wall_s
        qps = len(results) / actual_ws if actual_ws else 0.0

        w("| Queries | Wall Time | QPS | Memory |")
        w("|--------:|----------:|----:|-------:|")
        w(f"| {len(results)} | {actual_ws:.1f}s | {qps:.2f} | {peak_mb} MB |")
        w()

        suite_groups: Dict[str, List[QueryResult]] = defaultdict(list)
        for r in results:
            suite_groups[r.suite].append(r)
        SUITES = [s for s in ["HYBRID", "EXPLORE", "GUIDANCE", "ANALYTICS"] if suite_groups.get(s)]
        aggs = {s: _agg(suite_groups[s], s) for s in SUITES}

        w("| Suite | Routing | Correct/Total | Modes |")
        w("|-------|--------:|--------------:|-------|")
        for s in SUITES:
            a = aggs[s]
            bm = SUITE_BENCHMARKS.get(s, {}).get("routing")
            rt_pct = 100 * a["correct"] // a["n"] if a["n"] else 0
            grade = _grade(rt_pct, *bm) if bm else ""
            modes = "  ".join(f"{k}:{v}" for k, v in sorted(a["mode_ctr"].items(), key=lambda x: -x[1]))
            w(f"| {s} | {rt_pct}% {grade} | {a['correct']}/{a['n']} | {modes} |")
        w()

        w("| Suite | NDCG@5 | P@5 | R@5 | Coh@5 | p50 | p95 | p95 grade | Avg Cost |")
        w("|-------|-------:|----:|----:|------:|----:|----:|----------:|---------:|")
        for s in SUITES:
            a = aggs[s]
            lat_bm = SUITE_BENCHMARKS.get(s, {}).get("lat_p95_ms")
            if lat_bm and a["nl"]:
                warn_ms, good_ms, _dir, note = lat_bm
                p95_g = _lat_p95_grade(float(a["lat_p95"]), float(warn_ms), float(good_ms), note)
            else:
                p95_g = "n/a"
            w(
                f"| {s} | {a['ndcg5']:.3f} | {a['p5']:.3f} | {a['r5']:.3f} | {a['coh5']:.3f} | "
                f"{a['lat_p50']:.0f}ms | {a['lat_p95']:.0f}ms | {p95_g} | ${a['avg_cost_usd']:.6f} |"
            )
        w()

        w(
            "| Suite | SLA breach% | ≤4s% | ≤10s% | Timeouts | Errors | "
            "ZeroGuard | Fallback | Empty | Avg Results |"
        )
        w(
            "|-------|------------:|-----:|------:|---------:|-------:|"
            "----------:|---------:|------:|------------:|"
        )
        for s in SUITES:
            a = aggs[s]
            avg_res = a["total_results"] / a["nv"] if a["nv"] else 0.0
            w(
                f"| {s} | {pct(a['sla_b'], a['n'])} | {pct(a['le4'], a['nv'])} | "
                f"{pct(a['le10'], a['nv'])} | {a['timeouts']}/{a['n']} | "
                f"{a['errors']}/{a['n']} | {a['zero_guard']}/{a['n']} | "
                f"{a['explore_fb']}/{a['n']} | {a['zero_res']}/{a['n']} | {avg_res:.1f} |"
            )
        w()

        w("| Suite | Entity% | L0 | L1 | L2 | Fallback | Guard Recovery |")
        w("|-------|--------:|---:|---:|---:|---------:|---------------:|")
        for s in SUITES:
            a = aggs[s]
            ent_bm = SUITE_BENCHMARKS.get(s, {}).get("entity_rate")
            ent_grade = _grade(a["ent_pct_num"], *ent_bm) if (ent_bm and a["nv"]) else ""
            ent_pct = f"{pct(a['with_entity'], a['nv'])} {ent_grade}" if a["nv"] else "N/A"
            gf = a["guard_fired"]
            gr = f"{a['guard_ok_n']}/{len(gf)}={pct(a['guard_ok_n'], len(gf))}" if gf else "n/a"
            w(f"| {s} | {ent_pct} | {a['l0']}/{a['nv']} | {a['l1']}/{a['nv']} | {a['l2']}/{a['nv']} | {a['fb']}/{a['nv']} | {gr} |")
        w()
        # Surface any decision_tier string the eval taxonomy does not recognise —
        # catches future label drift instead of silently dropping it into the gap.
        _unknown_total: Counter = Counter()
        for s in SUITES:
            _unknown_total.update(aggs[s]["unknown_tiers"])
        if _unknown_total:
            w(f"> WARNING: Unrecognised decision_tier labels (not in L0/L1/L2/Fallback): "
              f"{dict(_unknown_total)} — update _ALL_KNOWN_TIERS in eval_test_search_queries.py")
            w()

        for suite in SUITES:
            w(f"## {suite}")
            w()
            sec_map: Dict[str, List[QueryResult]] = defaultdict(list)
            for r in suite_groups[suite]:
                sec_map[r.section].append(r)
            sections = list(sec_map.keys())
            saggs = {sec: _agg(sec_map[sec], suite) for sec in sections}

            w("| Section | Routing | Correct/Total | Modes |")
            w("|---------|--------:|--------------:|-------|")
            for sec in sections:
                a = saggs[sec]
                bm = SUITE_BENCHMARKS.get(suite, {}).get("routing")
                rt_pct = 100 * a["correct"] // a["n"] if a["n"] else 0
                grade = _grade(rt_pct, *bm) if bm else ""
                modes = "  ".join(f"{k}:{v}" for k, v in sorted(a["mode_ctr"].items(), key=lambda x: -x[1]))
                w(f"| {sec} | {rt_pct}% {grade} | {a['correct']}/{a['n']} | {modes} |")
            w()

            w("| Section | NDCG@5 | P@5 | R@5 | Coh@5 | p50 | p95 | Avg Cost |")
            w("|---------|-------:|----:|----:|------:|----:|----:|---------:|")
            for sec in sections:
                a = saggs[sec]
                w(f"| {sec} | {a['ndcg5']:.3f} | {a['p5']:.3f} | {a['r5']:.3f} | {a['coh5']:.3f} | {a['lat_p50']:.0f}ms | {a['lat_p95']:.0f}ms | ${a['avg_cost_usd']:.6f} |")
            w()

            w(
                "| Section | SLA breach% | ≤4s% | ≤10s% | Timeouts | Errors | "
                "ZeroGuard | Fallback | Empty | Avg Results |"
            )
            w(
                "|---------|------------:|-----:|------:|---------:|-------:|"
                "----------:|---------:|------:|------------:|"
            )
            for sec in sections:
                a = saggs[sec]
                avg_res = a["total_results"] / a["nv"] if a["nv"] else 0.0
                w(
                    f"| {sec} | {pct(a['sla_b'], a['n'])} | {pct(a['le4'], a['nv'])} | "
                    f"{pct(a['le10'], a['nv'])} | {a['timeouts']}/{a['n']} | "
                    f"{a['errors']}/{a['n']} | {a['zero_guard']}/{a['n']} | "
                    f"{a['explore_fb']}/{a['n']} | {a['zero_res']}/{a['n']} | {avg_res:.1f} |"
                )
            w()

            w("| Section | Entity% | L0 | L1 | L2 | Fallback | Guard Recovery |")
            w("|---------|--------:|---:|---:|---:|---------:|---------------:|")
            for sec in sections:
                a = saggs[sec]
                ent_bm = SUITE_BENCHMARKS.get(suite, {}).get("entity_rate")
                ent_grade = _grade(a["ent_pct_num"], *ent_bm) if (ent_bm and a["nv"]) else ""
                ent_pct = f"{pct(a['with_entity'], a['nv'])} {ent_grade}" if a["nv"] else "N/A"
                gf = a["guard_fired"]
                gr = f"{a['guard_ok_n']}/{len(gf)}={pct(a['guard_ok_n'], len(gf))}" if gf else "n/a"
                w(f"| {sec} | {ent_pct} | {a['l0']}/{a['nv']} | {a['l1']}/{a['nv']} | {a['l2']}/{a['nv']} | {a['fb']}/{a['nv']} | {gr} |")
            w()

        # Misroute table — suite expected mode ≠ observed answer_mode (excludes error/timeout).
        all_mis: List[QueryResult] = []
        for s in SUITES:
            all_mis.extend(aggs[s]["misroutes"])
        w("## Misroutes (suite ≠ answer_mode)")
        w()
        if not all_mis:
            w("_None_")
            w()
        else:
            w("| Suite | Section | Query | answer_mode | classified_intent | Expected modes |")
            w("|-------|---------|-------|-------------|-------------------|----------------|")
            for q in all_mis:
                expected = ", ".join(sorted(EXPECTED_MODE.get(q.suite, set())))
                q_short = (q.query[:80] + "…") if len(q.query) > 80 else q.query
                w(
                    f"| {q.suite} | {q.section} | {q_short} | {q.answer_mode} | "
                    f"{q.classified_intent or ''} | {expected} |"
                )
            w()

    w("# Evaluation Report")
    w()
    w(f"**API:** `{api}`")
    w()
    _render_api(results)

    with open(out_path, "w") as f:
        f.write("\n".join(lines))
    return out_path


# queries sheet columns (xlsx only). Top-k@5 still computed in search_one /
# aggregates via local_*_at; those dicts are not written as xlsx columns.
_XLSX_SCALAR_COLS = [
    "api_label", "suite", "section", "query",
    "latency_ms", "server_latency_ms", "latency_delta_ms",
    "answer_mode", "mode_correct", "classified_intent", "intent_confidence",
    "decision_tier", "multi_intent", "hard_filter_count", "total_candidates",
    "ndcg_at_5", "prec_at_5", "recall_at_5", "f1_at_5", "coh_at_5",
    "result_count", "avg_coherence",
    "timeout", "sla_breach", "failure_mode", "error",
    "entities_extracted_count", "decision_cost_usd",
    "stage_fusion_ms", "stage_eranker_ms",
    "zero_result_fired", "hard_gate_dropped",
]
_XLSX_LIST_COLS = [
    "retrieve_sources",
]
_XLSX_JSON_COLS = [
    "api_metrics", "ranked", "pipeline_stages",
]
_XLSX_QUERIES_COLS = (
    list(_XLSX_SCALAR_COLS) + list(_XLSX_LIST_COLS) + list(_XLSX_JSON_COLS)
)

_SUITE_ORDER = ["HYBRID", "EXPLORE", "GUIDANCE", "ANALYTICS"]


def _query_row_dict(r: QueryResult) -> Dict[str, Any]:
    """asdict + derived @5 / perf columns for xlsx/json."""
    d = asdict(r)
    d["ndcg_at_5"] = (r.local_ndcg_at or {}).get(5)
    d["prec_at_5"] = (r.local_prec_at or {}).get(5)
    d["recall_at_5"] = (r.local_recall_at or {}).get(5)
    d["f1_at_5"] = (r.local_f1_at or {}).get(5)
    d["coh_at_5"] = (r.local_coh_at or {}).get(5)
    d["result_count"] = r.ranked.count
    d["avg_coherence"] = r.ranked.avg_coherence
    stages = r.pipeline_stages or {}
    d["zero_result_fired"] = bool(stages.get("zero_result_fired")) if stages else None
    dropped = stages.get("hard_gate_dropped") if stages else None
    d["hard_gate_dropped"] = dropped
    if r.server_latency_ms is not None:
        d["latency_delta_ms"] = round(float(r.latency_ms) - float(r.server_latency_ms), 1)
    else:
        d["latency_delta_ms"] = None
    return d


def _metric_at(q: QueryResult, name: str, k: int, local_map: Dict[int, float]) -> float:
    """Prefer API ``retrieval_metrics`` at @k when present; else local proxy compute."""
    api_v = (q.api_metrics or {}).get(f"{name}@{k}")
    if api_v is not None:
        try:
            return float(api_v)
        except (TypeError, ValueError):
            pass
    return float((local_map or {}).get(k, 0.0))


# ZeroGuard column: ladder exhausted + rails empty (contracts.RankedResults).
_ZERO_GUARD_FAILURE_MODES = frozenset({
    "inventory_empty",
    "inventory_empty_under_filters",
})
# Fallback column: explore-rail / analytics failure fallback path (not empty-inventory).
_FALLBACK_FAILURE_MODES = frozenset({
    "explore_fallback_rail",
    "explore_fallback_timeout",
    "analytics_failure_fallback",
})


def _agg_for_report(qs: List[QueryResult], suite: str, sla_ms: int = _DEFAULT_SLA_MS) -> dict:
    """Shared suite/section aggregate used by markdown + xlsx summary sheets."""
    n = len(qs)
    valid = [q for q in qs if not q.error]
    nv = len(valid)
    lats = sorted(q.latency_ms for q in valid)
    nl = len(lats)
    suite_sla = _suite_sla_ms(suite, sla_ms)
    guard_fired = [q for q in qs if q.answer_mode == "explore_fallback" and suite != "EXPLORE"]
    guard_ok_n = sum(1 for q in guard_fired if q.mode_correct)
    lat_p50 = lats[nl // 2] if nl else 0.0
    lat_p90 = lats[min(int(nl * 0.90), nl - 1)] if nl else 0.0
    lat_p95 = lats[min(int(nl * 0.95), nl - 1)] if nl else 0.0
    lat_mean = statistics.mean(lats) if lats else 0.0
    le4 = sum(1 for q in valid if q.latency_ms <= LATENCY_LE_4S_MS)
    le10 = sum(1 for q in valid if q.latency_ms <= LATENCY_LE_10S_MS)
    misroutes = [
        q for q in qs
        if (not q.mode_correct) and (not q.error) and (not q.timeout)
    ]
    lat_bm = SUITE_BENCHMARKS.get(suite, {}).get("lat_p95_ms")
    if lat_bm and nl:
        warn_ms, good_ms, _dir, note = lat_bm
        p95_grade = _lat_p95_grade(float(lat_p95), float(warn_ms), float(good_ms), note)
    else:
        p95_grade = "n/a"
    costs = [q.decision_cost_usd for q in valid]
    # ZeroGuard ≠ Fallback (distinct signals; previously both used answer_mode).
    zero_guard_n = sum(
        1 for q in qs
        if (q.failure_mode or "") in _ZERO_GUARD_FAILURE_MODES and not q.error
    )
    explore_fb_n = sum(
        1 for q in qs
        if not q.error and (
            q.answer_mode == "explore_fallback"
            or (q.failure_mode or "") in _FALLBACK_FAILURE_MODES
        )
    )
    return dict(
        n=n, valid=valid, nv=nv, lats=lats, nl=nl, suite_sla=suite_sla,
        correct=sum(1 for q in qs if q.mode_correct),
        timeouts=sum(1 for q in qs if q.timeout),
        errors=sum(1 for q in qs if q.error),
        sla_b=sum(1 for q in qs if q.latency_ms > suite_sla and not q.error),
        zero_res=sum(1 for q in qs if q.ranked.zero_results and not q.error),
        explore_fb=explore_fb_n,
        zero_guard=zero_guard_n,
        total_results=sum(q.ranked.count for q in valid),
        mode_ctr=Counter(q.answer_mode or "unknown" for q in qs),
        ndcg5=(
            statistics.mean([
                _metric_at(q, "NDCG", 5, q.local_ndcg_at) for q in valid
            ]) if valid else 0.0
        ),
        p5=(
            statistics.mean([
                _metric_at(q, "Precision", 5, q.local_prec_at) for q in valid
            ]) if valid else 0.0
        ),
        r5=(
            statistics.mean([
                _metric_at(q, "Recall", 5, q.local_recall_at) for q in valid
            ]) if valid else 0.0
        ),
        f15=statistics.mean([q.local_f1_at.get(5, 0.0) for q in valid]) if valid else 0.0,
        coh5=(
            statistics.mean([
                _metric_at(q, "Coherence", 5, q.local_coh_at) for q in valid
            ]) if valid else 0.0
        ),
        lat_mean=lat_mean,
        lat_p50=lat_p50,
        lat_p90=lat_p90,
        lat_p95=lat_p95,
        p95_grade=p95_grade,
        le4=le4,
        le10=le10,
        misroutes=misroutes,
        misroute_n=len(misroutes),
        with_entity=sum(1 for q in valid if q.entities_extracted_count > 0),
        ent_pct_num=(100 * sum(1 for q in valid if q.entities_extracted_count > 0) // nv) if nv else 0,
        l0=sum(1 for q in valid if (q.decision_tier or "fallback") in _L0_TIERS),
        l1=sum(1 for q in valid if (q.decision_tier or "fallback") in _L1_TIERS),
        l2=sum(1 for q in valid if (q.decision_tier or "fallback") in _L2_TIERS),
        fb=sum(1 for q in valid if (q.decision_tier or "fallback") in _FALLBACK_TIERS),
        unknown_tiers=Counter(
            (q.decision_tier or "fallback") for q in valid
            if (q.decision_tier or "fallback") not in _ALL_KNOWN_TIERS
        ),
        guard_fired=guard_fired,
        guard_ok_n=guard_ok_n,
        avg_cost_usd=statistics.mean(costs) if costs else 0.0,
        total_cost_usd=sum(costs) if costs else 0.0,
        stages_present=sum(1 for q in valid if q.pipeline_stages),
    )


def _summary_row(label_key: str, label: str, a: dict) -> List[Any]:
    """One suite/section summary row for xlsx (no p95_grade / le_4s / le_10s)."""
    avg_res = a["total_results"] / a["nv"] if a["nv"] else 0.0
    rt_pct = (100 * a["correct"] // a["n"]) if a["n"] else 0
    gf = a["guard_fired"]
    gr = f"{a['guard_ok_n']}/{len(gf)}" if gf else "n/a"
    return [
        label,
        a["n"],
        rt_pct,
        f"{a['correct']}/{a['n']}",
        round(a["ndcg5"], 4),
        round(a["p5"], 4),
        round(a["r5"], 4),
        round(a["f15"], 4),
        round(a["coh5"], 4),
        round(a["lat_mean"], 1),
        round(a["lat_p50"], 1),
        round(a["lat_p90"], 1),
        round(a["lat_p95"], 1),
        round(a["avg_cost_usd"], 8),
        round(a["total_cost_usd"], 8),
        round(100.0 * a["sla_b"] / a["n"], 1) if a["n"] else None,
        a["timeouts"],
        a["errors"],
        a["zero_guard"],
        a["explore_fb"],
        a["zero_res"],
        round(avg_res, 2),
        a["ent_pct_num"] if a["nv"] else None,
        a["misroute_n"],
        a["l0"],
        a["l1"],
        a["l2"],
        a["fb"],
        gr,
        a["stages_present"],
        a["suite_sla"],
    ]


_SUMMARY_HEADERS = [
    "label", "n", "routing_pct", "correct_total",
    "ndcg_at_5", "prec_at_5", "recall_at_5", "f1_at_5", "coh_at_5",
    "lat_mean_ms", "lat_p50_ms", "lat_p90_ms", "lat_p95_ms",
    "avg_cost_usd", "total_cost_usd", "sla_breach_pct",
    "timeouts", "errors", "zero_guard", "fallback_explore", "empty",
    "avg_results", "entity_pct", "misroute_n",
    "L0", "L1", "L2", "tier_fallback",
    "guard_recovery", "stages_present", "suite_sla_ms",
]


# Formal stratified holdout (same contract as reground_filters_four_way).
DEFAULT_HOLDOUT_FRAC = 0.2
DEFAULT_HOLDOUT_SEED = 42


def _holdout_u01(query: str, *, seed: int, stratum: str) -> float:
    """Deterministic U[0,1) from md5(seed|stratum|query)."""
    digest = hashlib.md5(
        f"{int(seed)}|{stratum}|{query}".encode("utf-8"),
        usedforsecurity=False,
    ).hexdigest()
    return int(digest[:8], 16) / float(0xFFFFFFFF)


def assign_holdout_splits(
    rows: List[dict],
    *,
    frac: float,
    seed: int,
) -> Optional[Dict[str, Any]]:
    """Stamp each row with ``holdout_split`` = ``train``|``test`` (stratified by suite)."""
    if not (0.0 < float(frac) < 1.0):
        for r in rows:
            r.pop("holdout_split", None)
        return None
    frac_f = float(frac)
    seed_i = int(seed)
    by_suite: Dict[str, Dict[str, int]] = {}
    for r in rows:
        stratum = str(r.get("suite") or "all")
        q = str(r.get("query") or "")
        split = (
            "test"
            if _holdout_u01(q, seed=seed_i, stratum=stratum) < frac_f
            else "train"
        )
        r["holdout_split"] = split
        bucket = by_suite.setdefault(stratum, {"n": 0, "train": 0, "test": 0})
        bucket["n"] += 1
        bucket[split] += 1
    return {
        "seed": seed_i,
        "frac": frac_f,
        "method": "md5(seed|suite|query) < frac; stratified by suite",
        "unit": "query",
        "n_all": len(rows),
        "n_train": sum(1 for r in rows if r.get("holdout_split") == "train"),
        "n_test": sum(1 for r in rows if r.get("holdout_split") == "test"),
        "by_suite": by_suite,
    }


def _mean_flat(rows: List[dict], key: str) -> Optional[float]:
    vals = [float(r[key]) for r in rows if r.get(key) is not None]
    return round(statistics.mean(vals), 4) if vals else None


def _pctile(sorted_vals: List[float], p: float) -> Optional[float]:
    if not sorted_vals:
        return None
    n = len(sorted_vals)
    return round(sorted_vals[min(int(n * p), n - 1)], 1)


def _suite_metrics_from_dicts(rows: List[dict], suite: str) -> Dict[str, Any]:
    """Compact suite rollup from results JSON / row dicts (no QueryResult rebuild)."""
    n = len(rows)
    if n == 0:
        return {
            "suite": suite, "n": 0, "routing_pct": None, "correct_total": "0/0",
            "ndcg_at_5": None, "prec_at_5": None, "recall_at_5": None,
            "f1_at_5": None, "coh_at_5": None,
            "lat_mean_ms": None, "lat_p50_ms": None, "lat_p95_ms": None,
            "sla_breach_pct": None, "timeouts": 0, "errors": 0,
            "avg_cost_usd": None, "misroute_n": 0,
        }
    valid = [r for r in rows if not r.get("error")]
    correct = sum(1 for r in rows if r.get("mode_correct"))
    suite_sla = _suite_sla_ms(suite if suite != "ALL" else "HYBRID", _DEFAULT_SLA_MS)
    if suite == "ALL":
        # Per-row SLA: use that row's suite threshold.
        sla_b = sum(
            1 for r in rows
            if (not r.get("error"))
            and float(r.get("latency_ms") or 0) > _suite_sla_ms(
                str(r.get("suite") or "HYBRID"), _DEFAULT_SLA_MS,
            )
        )
    else:
        sla_b = sum(
            1 for r in rows
            if (not r.get("error"))
            and float(r.get("latency_ms") or 0) > suite_sla
        )
    lats = sorted(
        float(r["latency_ms"]) for r in valid if r.get("latency_ms") is not None
    )
    misroute_n = sum(
        1 for r in rows
        if (not r.get("mode_correct"))
        and (not r.get("error"))
        and (not r.get("timeout"))
    )
    return {
        "suite": suite,
        "n": n,
        "routing_pct": round(100.0 * correct / n, 1),
        "correct_total": f"{correct}/{n}",
        "ndcg_at_5": _mean_flat(valid, "ndcg_at_5"),
        "prec_at_5": _mean_flat(valid, "prec_at_5"),
        "recall_at_5": _mean_flat(valid, "recall_at_5"),
        "f1_at_5": _mean_flat(valid, "f1_at_5"),
        "coh_at_5": _mean_flat(valid, "coh_at_5"),
        "lat_mean_ms": round(statistics.mean(lats), 1) if lats else None,
        "lat_p50_ms": _pctile(lats, 0.50),
        "lat_p95_ms": _pctile(lats, 0.95),
        "sla_breach_pct": round(100.0 * sla_b / n, 1),
        "timeouts": sum(1 for r in rows if r.get("timeout")),
        "errors": sum(1 for r in rows if r.get("error")),
        "avg_cost_usd": _mean_flat(valid, "decision_cost_usd"),
        "misroute_n": misroute_n,
    }


def build_eval_holdout_report(
    rows: List[dict],
    *,
    frac: float = DEFAULT_HOLDOUT_FRAC,
    seed: int = DEFAULT_HOLDOUT_SEED,
) -> Optional[Dict[str, Any]]:
    """Assign holdout splits and compute all/train/test suite metrics.

    Full-search eval only (live ``/search`` routing + retrieval quality).
    Not arm pairwise; not LLMJ / regex comparison.
    """
    meta = assign_holdout_splits(rows, frac=frac, seed=seed)
    if meta is None:
        return None
    splits: Dict[str, List[dict]] = {
        "all": rows,
        "train": [r for r in rows if r.get("holdout_split") == "train"],
        "test": [r for r in rows if r.get("holdout_split") == "test"],
    }
    by_split: Dict[str, Any] = {}
    focus: List[Dict[str, Any]] = []
    for split_name, subset in splits.items():
        suites_present = sorted({str(r.get("suite") or "") for r in subset if r.get("suite")})
        suite_blocks = {
            suite: _suite_metrics_from_dicts(
                [r for r in subset if r.get("suite") == suite], suite,
            )
            for suite in suites_present
        }
        overall = _suite_metrics_from_dicts(subset, "ALL")
        by_split[split_name] = {"ALL": overall, "by_suite": suite_blocks}
        focus.append({
            "split": split_name,
            **{k: overall[k] for k in (
                "n", "routing_pct", "correct_total",
                "ndcg_at_5", "prec_at_5", "recall_at_5", "coh_at_5",
                "lat_p50_ms", "lat_p95_ms", "sla_breach_pct",
                "timeouts", "errors", "misroute_n", "avg_cost_usd",
            )},
        })
        for suite, m in suite_blocks.items():
            focus.append({
                "split": split_name,
                "suite": suite,
                **{k: m[k] for k in (
                    "n", "routing_pct", "correct_total",
                    "ndcg_at_5", "prec_at_5", "recall_at_5", "coh_at_5",
                    "lat_p50_ms", "lat_p95_ms", "sla_breach_pct",
                    "timeouts", "errors", "misroute_n", "avg_cost_usd",
                )},
            })
    # Ensure ALL focus rows have suite key for table uniformity.
    for row in focus:
        row.setdefault("suite", "ALL")
    return {**meta, "splits": by_split, "focus": focus}


def _print_holdout_focus(holdout: Dict[str, Any]) -> None:
    print(
        f"HOLDOUT seed={holdout['seed']} frac={holdout['frac']} "
        f"train={holdout['n_train']} test={holdout['n_test']}"
    )
    for row in holdout.get("focus") or []:
        if row.get("suite") != "ALL":
            continue
        print(
            f"HOLDOUT_FOCUS split={row['split']} suite=ALL "
            f"n={row['n']} routing={row['routing_pct']}% "
            f"ndcg@5={row['ndcg_at_5']} p95={row['lat_p95_ms']}ms "
            f"sla_breach={row['sla_breach_pct']}%"
        )


def _write_eval_sheet_jsons(
    rows: List[dict],
    out_dir: Path,
    stamp: str,
    *,
    holdout: Optional[Dict[str, Any]],
) -> Dict[str, Path]:
    """Sheet-wise JSON mirroring xlsx (holdout lives under suite_summary)."""
    paths: Dict[str, Path] = {}
    suites_present = sorted({str(r.get("suite") or "") for r in rows if r.get("suite")})
    suite_blocks = {
        suite: _suite_metrics_from_dicts(
            [r for r in rows if r.get("suite") == suite], suite,
        )
        for suite in suites_present
    }
    suite_payload: Dict[str, Any] = {
        "ALL": _suite_metrics_from_dicts(rows, "ALL"),
        "by_suite": suite_blocks,
        "holdout": holdout,
    }
    paths["suite_summary"] = out_dir / f"suite_summary_{stamp}.json"
    paths["suite_summary"].write_text(json.dumps(suite_payload, indent=2))

    section_blocks: Dict[str, Any] = {}
    for r in rows:
        suite = str(r.get("suite") or "")
        section = str(r.get("section") or "")
        if not suite:
            continue
        key = f"{suite} / {section}" if section else suite
        section_blocks.setdefault(key, []).append(r)
    section_payload = {
        key: _suite_metrics_from_dicts(subset, key)
        for key, subset in sorted(section_blocks.items())
    }
    paths["section_summary"] = out_dir / f"section_summary_{stamp}.json"
    paths["section_summary"].write_text(json.dumps(section_payload, indent=2))

    fail_rows = [
        r for r in rows
        if r.get("error") or r.get("timeout") or r.get("sla_breach")
        or r.get("mode_correct") is False
    ]
    paths["fail_only"] = out_dir / f"fail_only_{stamp}.json"
    paths["fail_only"].write_text(json.dumps(fail_rows, indent=2))

    mis_rows = [
        r for r in rows
        if r.get("mode_correct") is False
        and not r.get("error")
        and not r.get("timeout")
    ]
    paths["misroutes"] = out_dir / f"misroutes_{stamp}.json"
    paths["misroutes"].write_text(json.dumps(mis_rows, indent=2))
    return paths


def _write_json_artifacts(
    results: List[QueryResult],
    out_dir: Path,
    stamp: str,
    *,
    holdout_frac: float = DEFAULT_HOLDOUT_FRAC,
    holdout_seed: int = DEFAULT_HOLDOUT_SEED,
) -> Tuple[Path, Dict[str, Path], Optional[Dict[str, Any]]]:
    """Write ``results_{stamp}.json`` (queries) + sheet-wise JSON per xlsx sheet.

    Holdout is nested under ``suite_summary_*.json`` and appended on the
    ``suite_summary`` xlsx sheet — no separate holdout_*.json.
    """
    out_dir.mkdir(parents=True, exist_ok=True)
    rows = [_query_row_dict(r) for r in results]
    holdout = build_eval_holdout_report(
        rows, frac=holdout_frac, seed=holdout_seed,
    )
    # Strip split stamps so results_*.json schema stays identical to pre-holdout.
    for r in rows:
        r.pop("holdout_split", None)
    results_path = out_dir / f"results_{stamp}.json"
    results_path.write_text(json.dumps(rows, indent=2))
    n_fail = sum(1 for r in results if _is_fail_row(r))
    print(f"JSON: {results_path}  fail_rows={n_fail}/{len(rows)}")
    sheet_paths = _write_eval_sheet_jsons(
        rows, out_dir, stamp, holdout=holdout,
    )
    for name, path in sheet_paths.items():
        print(f"SHEET_JSON {name} -> {path}")
    if holdout is not None:
        _print_holdout_focus(holdout)
    else:
        print(f"HOLDOUT disabled (holdout_frac={holdout_frac})")
    return results_path, sheet_paths, holdout


def _write_xlsx_report(
    results: List[QueryResult],
    out_dir: Path,
    stamp: str,
    sla_ms: int = _DEFAULT_SLA_MS,
    holdout: Optional[Dict[str, Any]] = None,
) -> Path:
    """Write multi-sheet `eval_report_{stamp}.xlsx`.

    Sheets: ``queries`` (per-query, keeps @5 + cost + guards + new stage/band cols),
    ``suite_summary`` (full-corpus rollup + appended HOLDOUT train/test block),
    ``section_summary``, ``fail_only``, ``misroutes``, ``description``.
    No separate holdout sheet.
    """
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"eval_report_{stamp}.xlsx"
    wb = Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    blue_fill = PatternFill("solid", fgColor="1F4E79")
    green_fill = PatternFill("solid", fgColor="1B5E20")
    orange_fill = PatternFill("solid", fgColor="BF360C")
    purple_fill = PatternFill("solid", fgColor="4A148C")
    center = Alignment(horizontal="center")

    def _autowidth(ws) -> None:
        for col_cells in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col_cells), default=10)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 60)

    def _write_header(ws, cols: List[str], fill_for) -> None:
        for c, name in enumerate(cols, 1):
            cell = ws.cell(row=1, column=c, value=name)
            cell.font = header_font
            cell.fill = fill_for(name)
            cell.alignment = center
        ws.freeze_panes = "A2"

    # ── queries ──────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "queries"
    cols = list(_XLSX_QUERIES_COLS)
    _at5_cols = {"ndcg_at_5", "prec_at_5", "recall_at_5", "f1_at_5", "coh_at_5"}
    _perf_cols = {
        "latency_ms", "server_latency_ms", "latency_delta_ms",
        "stage_fusion_ms", "stage_eranker_ms", "zero_result_fired", "hard_gate_dropped",
    }

    def _q_fill(name: str):
        if name in _XLSX_JSON_COLS:
            return orange_fill
        if name in _XLSX_LIST_COLS:
            return green_fill
        if name in _at5_cols:
            return purple_fill
        if name in _perf_cols:
            return orange_fill
        return blue_fill

    _write_header(ws, cols, _q_fill)
    for ri, r in enumerate(results, 2):
        d = _query_row_dict(r)
        for ci, name in enumerate(cols, 1):
            if name in _XLSX_LIST_COLS:
                val = ", ".join(d.get(name) or [])
            elif name in _XLSX_JSON_COLS:
                val = json.dumps(d.get(name), sort_keys=True)
            else:
                val = d.get(name, "")
            ws.cell(row=ri, column=ci, value=val)
    _autowidth(ws)

    # ── suite_summary / section_summary ──────────────────────────────────────
    suite_groups: Dict[str, List[QueryResult]] = defaultdict(list)
    for r in results:
        suite_groups[r.suite].append(r)
    suites = [s for s in _SUITE_ORDER if suite_groups.get(s)]

    def _fill_summary_sheet(ws, headers: List[str], data_rows: List[list]) -> None:
        _write_header(ws, headers, lambda _n: blue_fill)
        for ri, row in enumerate(data_rows, 2):
            for ci, val in enumerate(row, 1):
                ws.cell(row=ri, column=ci, value=val)
        _autowidth(ws)

    suite_rows = []
    section_rows = []
    all_mis: List[QueryResult] = []
    for s in suites:
        a = _agg_for_report(suite_groups[s], s, sla_ms)
        suite_rows.append(_summary_row("suite", s, a))
        all_mis.extend(a["misroutes"])
        sec_map: Dict[str, List[QueryResult]] = defaultdict(list)
        for r in suite_groups[s]:
            sec_map[r.section].append(r)
        for sec, qs in sec_map.items():
            sa = _agg_for_report(qs, s, sla_ms)
            section_rows.append(_summary_row("section", f"{s} / {sec}", sa))

    suite_headers = list(_SUMMARY_HEADERS)
    suite_headers[0] = "suite"
    section_headers = list(_SUMMARY_HEADERS)
    section_headers[0] = "section"

    ws_s = wb.create_sheet("suite_summary")
    _fill_summary_sheet(ws_s, suite_headers, suite_rows)
    # Holdout train/test — appended on suite_summary (no new sheet).
    if holdout:
        h_cols = [
            "split", "suite", "n", "routing_pct", "correct_total",
            "ndcg_at_5", "prec_at_5", "recall_at_5", "coh_at_5",
            "lat_p50_ms", "lat_p95_ms", "sla_breach_pct",
            "timeouts", "errors", "misroute_n", "avg_cost_usd",
        ]
        start = int(ws_s.max_row or 1) + 2
        ws_s.cell(
            row=start, column=1,
            value=(
                f"HOLDOUT train/test (seed={holdout.get('seed')} "
                f"frac={holdout.get('frac')} train={holdout.get('n_train')} "
                f"test={holdout.get('n_test')}; {holdout.get('method')})"
            ),
        )
        ws_s.cell(row=start, column=1).font = Font(bold=True)
        for ci, name in enumerate(h_cols, 1):
            cell = ws_s.cell(row=start + 1, column=ci, value=name)
            cell.font = header_font
            cell.fill = green_fill
        for ri, row in enumerate(holdout.get("focus") or [], start + 2):
            for ci, name in enumerate(h_cols, 1):
                ws_s.cell(row=ri, column=ci, value=row.get(name))
        _autowidth(ws_s)
    ws_sec = wb.create_sheet("section_summary")
    _fill_summary_sheet(ws_sec, section_headers, section_rows)

    # ── fail_only ────────────────────────────────────────────────────────────
    fail_results = [r for r in results if _is_fail_row(r)]
    ws_f = wb.create_sheet("fail_only")
    fail_cols = [
        "suite", "section", "query", "latency_ms", "answer_mode", "mode_correct",
        "classified_intent", "decision_tier",
        "ndcg_at_5", "prec_at_5", "recall_at_5", "coh_at_5",
        "timeout", "sla_breach",
        "failure_mode", "error", "decision_cost_usd", "server_latency_ms",
        "zero_results", "backends_active", "retrieve_sources",
    ]
    _write_header(ws_f, fail_cols, lambda _n: orange_fill)
    for ri, r in enumerate(fail_results, 2):
        d = _query_row_dict(r)
        d["zero_results"] = r.ranked.zero_results
        for ci, name in enumerate(fail_cols, 1):
            if name in ("backends_active", "retrieve_sources"):
                val = ", ".join(d.get(name) or [])
            else:
                val = d.get(name, "")
            ws_f.cell(row=ri, column=ci, value=val)
    _autowidth(ws_f)

    # ── misroutes ────────────────────────────────────────────────────────────
    ws_m = wb.create_sheet("misroutes")
    mis_cols = [
        "suite", "section", "query", "answer_mode", "classified_intent",
        "expected_modes", "latency_ms", "decision_tier",
        "ndcg_at_5", "prec_at_5", "recall_at_5", "coh_at_5", "decision_cost_usd",
    ]
    _write_header(ws_m, mis_cols, lambda _n: purple_fill)
    for ri, r in enumerate(all_mis, 2):
        d = _query_row_dict(r)
        d["expected_modes"] = ", ".join(sorted(EXPECTED_MODE.get(r.suite, set())))
        for ci, name in enumerate(mis_cols, 1):
            ws_m.cell(row=ri, column=ci, value=d.get(name, ""))
    _autowidth(ws_m)

    # ── description ──────────────────────────────────────────────────────────
    ws_d = wb.create_sheet("description")
    ws_d["A1"] = "Sheet / column"
    ws_d["B1"] = "Meaning"
    ws_d["A1"].font = header_font
    ws_d["B1"].font = header_font
    ws_d["A1"].fill = blue_fill
    ws_d["B1"].fill = blue_fill
    desc_rows = [
        ("queries",
         "One row per query: @5 quality scalars, client/server/delta latency, cost, "
         "result_count/avg_coherence, zero_result_fired/hard_gate_dropped, stages JSON. "
         "Dropped from xlsx: started_at, API ndcg/coherence/recall/precision, local_*_at dicts, "
         "filter_names/backends_active/type_id_entities, under_4s/under_10s."),
        ("ndcg_at_5 / prec_at_5 / recall_at_5 / f1_at_5 / coh_at_5",
         "Prefer retrieval_metrics NDCG@5/Precision@5/... when present; else local proxy "
         f"from ranked_results[].coherence_score via _ndcg_at/_prf_at "
         f"(rel_threshold={_REL_THRESHOLD} from retrieval.metrics.relevance_threshold; "
         "score>=threshold = relevant; no gold labels). F1 always local."),
        ("zero_guard / fallback_explore / empty",
         "Distinct: zero_guard=failure_mode in {inventory_empty,inventory_empty_under_filters}; "
         "fallback_explore=answer_mode==explore_fallback OR failure_mode in "
         "{explore_fallback_rail,explore_fallback_timeout,analytics_failure_fallback}; "
         "empty=ranked_results empty (zero_results)."),
        ("latency_ms / server_latency_ms / latency_delta_ms",
         "Client HTTP wall; body.latency_ms; delta = client - server (None if server missing). "
         "sla_breach uses client latency_ms vs suite SLA."),
        ("result_count / avg_coherence",
         "From ranked_results interpretation (count + mean coherence_score)."),
        ("zero_result_fired / hard_gate_dropped",
         "From pipeline_trace.stages when present; None if stages empty."),
        ("decision_cost_usd",
         "Copied from query_intelligence.decision_cost_usd (QI/LLM path spend)."),
        ("entities_extracted_count",
         "len(hard_filters from identified∪applied_filters; keyword_* excluded); "
         "falls back to keyword term count when no hard chips."),
        ("suite_summary / section_summary",
         "Rollups: routing, @5 (incl f1), lat mean/p50/p90/p95, avg+total cost, sla_breach%, "
         "timeouts/errors/zero_guard/fallback/empty/avg_results, entity%, misroute_n, "
         "L0/L1/L2, guard recovery. No p95_grade / le_4s_pct / le_10s_pct on these sheets."),
        ("suite_summary (HOLDOUT block)",
         "Appended below full-corpus suite rows: stratified train/test metrics "
         "(seed/frac). Same numbers as suite_summary_*.json[\"holdout\"]. No separate sheet."),
        ("fail_only", "error OR timeout OR sla_breach OR mode_incorrect."),
        ("misroutes", "mode_correct=False excluding error/timeout; expected modes from EXPECTED_MODE."),
        ("results_*.json", "queries sheet — per-query dump (analysis-only source; no holdout columns)."),
        ("suite_summary_*.json", "suite rollup + holdout nested under \"holdout\" key."),
        ("section_summary_*.json / fail_only_*.json / misroutes_*.json",
         "Sheet-wise JSON mirroring the xlsx sheets of the same name."),
    ]
    for i, (k, v) in enumerate(desc_rows, 2):
        ws_d.cell(row=i, column=1, value=k)
        ws_d.cell(row=i, column=2, value=v)
    _autowidth(ws_d)

    wb.save(out_path)
    print(f"XLSX: {out_path}")
    return out_path


def _resolve_latest_eval_results(out_dir: Path, explicit: Optional[Path]) -> Path:
    if explicit is not None:
        path = explicit.resolve()
        if not path.is_file():
            raise SystemExit(f"--results not found: {path}")
        return path
    candidates = sorted(out_dir.glob("results_*.json"), key=lambda p: p.stat().st_mtime)
    if not candidates:
        raise SystemExit(f"--analysis-only needs --results PATH or results_*.json under {out_dir}")
    return candidates[-1]


def _append_holdout_to_suite_summary_sheet(
    xlsx_path: Path, holdout: Dict[str, Any],
) -> None:
    """Append HOLDOUT block to existing suite_summary sheet (no new sheet)."""
    wb = load_workbook(xlsx_path)
    if "suite_summary" not in wb.sheetnames:
        wb.close()
        raise SystemExit(f"suite_summary sheet missing in {xlsx_path}")
    ws = wb["suite_summary"]
    # Drop any prior HOLDOUT block so re-runs stay idempotent.
    cut_at: Optional[int] = None
    for r in range(1, (ws.max_row or 1) + 1):
        val = ws.cell(row=r, column=1).value
        if isinstance(val, str) and val.startswith("HOLDOUT train/test"):
            cut_at = r
            break
    if cut_at is not None:
        ws.delete_rows(cut_at, (ws.max_row or cut_at) - cut_at + 1)
    h_cols = [
        "split", "suite", "n", "routing_pct", "correct_total",
        "ndcg_at_5", "prec_at_5", "recall_at_5", "coh_at_5",
        "lat_p50_ms", "lat_p95_ms", "sla_breach_pct",
        "timeouts", "errors", "misroute_n", "avg_cost_usd",
    ]
    start = int(ws.max_row or 1) + 2
    ws.cell(
        row=start, column=1,
        value=(
            f"HOLDOUT train/test (seed={holdout.get('seed')} "
            f"frac={holdout.get('frac')} train={holdout.get('n_train')} "
            f"test={holdout.get('n_test')}; {holdout.get('method')})"
        ),
    )
    ws.cell(row=start, column=1).font = Font(bold=True)
    header_font = Font(bold=True, color="FFFFFF")
    green_fill = PatternFill("solid", fgColor="1B5E20")
    for ci, name in enumerate(h_cols, 1):
        cell = ws.cell(row=start + 1, column=ci, value=name)
        cell.font = header_font
        cell.fill = green_fill
    for ri, row in enumerate(holdout.get("focus") or [], start + 2):
        for ci, name in enumerate(h_cols, 1):
            ws.cell(row=ri, column=ci, value=row.get(name))
    wb.save(xlsx_path)
    wb.close()
    print(f"HOLDOUT appended to suite_summary in {xlsx_path}")


def _analysis_only_from_results(
    results_path: Path,
    out_dir: Path,
    *,
    holdout_frac: float,
    holdout_seed: int,
) -> Path:
    """Rebuild sheet-wise JSON + append HOLDOUT onto existing eval_report xlsx."""
    rows = json.loads(results_path.read_text())
    if not isinstance(rows, list) or not rows:
        raise SystemExit(f"empty/invalid results JSON: {results_path}")
    stamp = str(int(time.time()))
    holdout = build_eval_holdout_report(rows, frac=holdout_frac, seed=holdout_seed)
    for r in rows:
        r.pop("holdout_split", None)
    out_dir.mkdir(parents=True, exist_ok=True)
    print(f"ANALYSIS_ONLY results_in={results_path} n={len(rows)} (results file left unchanged)")
    sheet_paths = _write_eval_sheet_jsons(
        rows, out_dir, stamp, holdout=holdout,
    )
    for name, path in sheet_paths.items():
        print(f"SHEET_JSON {name} -> {path}")
    if holdout is not None:
        _print_holdout_focus(holdout)
        reports = sorted(
            out_dir.glob("eval_report_*.xlsx"), key=lambda p: p.stat().st_mtime,
        )
        if reports:
            _append_holdout_to_suite_summary_sheet(reports[-1], holdout)
        else:
            print("HOLDOUT_XLSX skipped — no eval_report_*.xlsx to append into")
    else:
        print(f"HOLDOUT disabled (holdout_frac={holdout_frac})")
    return sheet_paths.get("suite_summary") or results_path


async def main() -> Tuple[Path, Path]:
    argv = sys.argv[1:]
    if not argv or argv[0] in ("-h", "--help"):
        sys.stderr.write("Usage: python -m semantic_search.offline_harness.eval_test_search_queries <url> [--suite S1,S2]\n")
        sys.stderr.write("       python -m semantic_search.offline_harness.eval_test_search_queries --api <url> [--suite S1,S2]\n")
        sys.stderr.write("       python -m semantic_search.offline_harness.eval_test_search_queries --analysis-only [--results PATH]\n")
        sys.stderr.write("  --api            API base URL (e.g. http://localhost:8085)\n")
        sys.stderr.write("  --suite          Comma-separated suites (HYBRID,EXPLORE,GUIDANCE,ANALYTICS)\n")
        sys.stderr.write("  --analysis-only  Rebuild holdout from results_*.json (no HTTP)\n")
        sys.stderr.write("  --results        results_*.json path (with --analysis-only)\n")
        sys.stderr.write(
            f"  --holdout-frac   Train/test holdout fraction (default {DEFAULT_HOLDOUT_FRAC}; 0 disables)\n"
        )
        sys.stderr.write(f"  --holdout-seed   Holdout seed (default {DEFAULT_HOLDOUT_SEED})\n")
        sys.exit(0 if argv else 1)

    suite_filter: Optional[Set[str]] = None
    api: Optional[str] = None
    analysis_only = False
    results_arg: Optional[Path] = None
    holdout_frac = DEFAULT_HOLDOUT_FRAC
    holdout_seed = DEFAULT_HOLDOUT_SEED
    i = 0
    while i < len(argv):
        if argv[i] == "--suite" and i + 1 < len(argv):
            suite_filter = set(argv[i + 1].upper().split(","))
            i += 2
        elif argv[i] in ("--api", "--api-a") and i + 1 < len(argv):
            api = argv[i + 1].rstrip("/")
            i += 2
        elif argv[i] == "--analysis-only":
            analysis_only = True
            i += 1
        elif argv[i] == "--results" and i + 1 < len(argv):
            results_arg = Path(argv[i + 1])
            i += 2
        elif argv[i] == "--holdout-frac" and i + 1 < len(argv):
            holdout_frac = float(argv[i + 1])
            i += 2
        elif argv[i] == "--holdout-seed" and i + 1 < len(argv):
            holdout_seed = int(argv[i + 1])
            i += 2
        elif not argv[i].startswith("--"):
            api = argv[i].rstrip("/")
            i += 1
        else:
            i += 1

    if analysis_only:
        results_path = _resolve_latest_eval_results(DEFAULT_OUTPUT_DIR, results_arg)
        out = _analysis_only_from_results(
            results_path, DEFAULT_OUTPUT_DIR,
            holdout_frac=holdout_frac, holdout_seed=holdout_seed,
        )
        return out, out

    if not api:
        sys.stderr.write("Error: API URL required (positional or --api <url>)\n")
        sys.exit(1)
    # SSRF guard (CLI): localhost / *.internal / *.local only — same allowlist as
    # query_eval_runner. HTTP harness route applies a stricter loopback-only check.
    from semantic_search.eval.query_eval_runner import (  # noqa: E402
        _validate_endpoint_url,
    )
    try:
        _validate_endpoint_url(api if "://" in api else f"http://{api}")
    except ValueError as exc:
        sys.stderr.write(f"Error: invalid --api URL ({exc})\n")
        sys.exit(1)

    queries = list(QUERY_BANK)
    if suite_filter:
        queries = [(s, sec, q) for s, sec, q in queries if s in suite_filter]

    tracemalloc.start()
    t_wall_start = time.perf_counter()

    try:
        results = await run_suite_queries(
            "API", api, queries, _CONCURRENCY, _TOP_K, _DEFAULT_SLA_MS, _TIMEOUT_S, _REL_THRESHOLD
        )
    except httpx.ConnectError as e:
        sys.stderr.write(f"Fatal: {api} unreachable — {e}\nAborting eval.\n")
        sys.exit(1)

    wall_s = time.perf_counter() - t_wall_start
    _, peak_bytes = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    peak_mb = round(peak_bytes / 1_048_576, 2)
    total_queries_run = len(results)

    stamp = str(int(time.time()))
    # Write results + sheet JSON first so artifacts exist even if report writers fail later.
    results_json, sheet_paths, holdout = _write_json_artifacts(
        results, DEFAULT_OUTPUT_DIR, stamp,
        holdout_frac=holdout_frac, holdout_seed=holdout_seed,
    )
    md_path = _write_markdown_report(
        results=results,
        api=api,
        wall_s=wall_s,
        peak_mb=peak_mb,
        total_queries_run=total_queries_run,
        sla_ms=_DEFAULT_SLA_MS,
        out_dir=DEFAULT_OUTPUT_DIR,
        stamp=stamp,
    )
    xlsx_path = _write_xlsx_report(
        results, DEFAULT_OUTPUT_DIR, stamp,
        sla_ms=_DEFAULT_SLA_MS, holdout=holdout,
    )
    sheet_lines = "".join(f"SHEET_JSON {n} -> {p}\n" for n, p in sheet_paths.items())
    sys.stderr.write(
        f"Report -> {md_path}\nXLSX -> {xlsx_path}\n"
        f"JSON -> {results_json}\n{sheet_lines}"
    )
    return md_path, xlsx_path


if __name__ == "__main__":
    asyncio.run(main())
